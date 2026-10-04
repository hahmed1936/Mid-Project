"""Build the discussion deck for the Airline Passenger Satisfaction mid project (python-pptx, native charts)."""
import pandas as pd
import numpy as np
from scipy import stats
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_LABEL_POSITION

from pathlib import Path
PROJECT = str(Path(__file__).parent)
OUT = PROJECT + r"\Airline_Satisfaction_Presentation.pptx"

# ---------------------------------------------------------------- data
df = pd.read_csv(PROJECT + r"\Data_Cleaned.csv")
raw = pd.read_csv(PROJECT + r"\Data.csv")
SERVICES = list(df.loc[:, "inflight_wifi_service":"cleanliness"].columns)


def rate(col, order=None):
    s = df.groupby(col)["is_satisfied"].mean() * 100
    return s.reindex(order) if order else s


def cramers_v(col):
    t = pd.crosstab(df[col], df["satisfaction"])
    chi2 = stats.chi2_contingency(t)[0]
    return np.sqrt(chi2 / (t.values.sum() * (min(t.shape) - 1)))


# ---------------------------------------------------------------- design
NAVY = RGBColor(0x0B, 0x1F, 0x3A)
NAVY_2 = RGBColor(0x15, 0x30, 0x55)
SKY = RGBColor(0x2A, 0x78, 0xD6)
ORANGE = RGBColor(0xEB, 0x68, 0x34)
ICE = RGBColor(0xE8, 0xF1, 0xFB)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
INK = RGBColor(0x1B, 0x1F, 0x24)
MUTED = RGBColor(0x5B, 0x63, 0x6E)
LIGHT_MUTED = RGBColor(0xB8, 0xC7, 0xDB)
GRID = RGBColor(0xDD, 0xE3, 0xEA)
FONT = "Calibri"
HEAD_FONT = "Cambria"

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
BLANK = prs.slide_layouts[6]
SW, SH = 13.333, 7.5


def bg(slide, color):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def text(slide, x, y, w, h, content, size=16, color=INK, bold=False, font=FONT, align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP, italic=False, name=None):
    """Add a text box. `content` may be a string or a list of (text, options) paragraphs."""
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    if name:
        box.name = name
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    paragraphs = content if isinstance(content, list) else [(content, {})]
    for i, (txt, opts) in enumerate(paragraphs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = opts.get("align", align)
        p.space_after = Pt(opts.get("space_after", 0))
        runs = txt if isinstance(txt, list) else [(txt, {})]
        for rtxt, ropts in runs:
            r = p.add_run()
            r.text = rtxt
            f = r.font
            f.name = ropts.get("font", opts.get("font", font))
            f.size = Pt(ropts.get("size", opts.get("size", size)))
            f.bold = ropts.get("bold", opts.get("bold", bold))
            f.italic = ropts.get("italic", opts.get("italic", italic))
            f.color.rgb = ropts.get("color", opts.get("color", color))
    return box


def shape(slide, kind, x, y, w, h, fill, line=None, name=None, shadow=False):
    s = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Pt(1)
    if not shadow:
        s.shadow.inherit = False
    if kind == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = 0.08
    if name:
        s.name = name
    return s


def badge(slide, x, y, label, fill=SKY, size=0.55, font_size=16, text_color=WHITE):
    """The deck's motif: a numbered circle badge."""
    c = shape(slide, MSO_SHAPE.OVAL, x, y, size, size, fill)
    tf = c.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    r = p.add_run()
    r.text = label
    r.font.size, r.font.bold, r.font.color.rgb, r.font.name = Pt(font_size), True, text_color, FONT
    return c


slide_no = [0]


def content_slide(title, kicker=None):
    """Light content slide with a consistent title position and a footer."""
    s = prs.slides.add_slide(BLANK)
    bg(s, WHITE)
    slide_no[0] += 1
    if kicker:
        text(s, 0.6, 0.38, 10, 0.3, kicker.upper(), size=12, color=SKY, bold=True)
    text(s, 0.6, 0.68, 12.1, 0.8, title, size=32, color=NAVY, bold=True, font=HEAD_FONT, name="Title")
    text(s, 0.6, 7.0, 8, 0.3, "Airline Passenger Satisfaction  ·  Mid Project", size=10, color=MUTED)
    text(s, 12.2, 7.0, 0.55, 0.3, str(slide_no[0] + 1), size=10, color=MUTED, align=PP_ALIGN.RIGHT)
    return s


def dark_slide():
    s = prs.slides.add_slide(BLANK)
    bg(s, NAVY)
    slide_no[0] += 1
    return s


def style_chart(chart, legend=False, value_axis_max=None, number_format='0"%"', font_size=12, gridlines=True):
    chart.has_title = False
    chart.font.name = FONT
    chart.font.size = Pt(font_size)
    chart.font.color.rgb = MUTED
    chart.has_legend = legend
    if legend:
        chart.legend.position = XL_LEGEND_POSITION.TOP
        chart.legend.include_in_layout = False
        chart.legend.font.size = Pt(12)
        chart.legend.font.color.rgb = INK
    try:
        va = chart.value_axis
        va.has_major_gridlines = gridlines
        if gridlines:
            va.major_gridlines.format.line.color.rgb = GRID
        va.format.line.fill.background()
        va.tick_labels.font.size = Pt(11)
        va.tick_labels.number_format = number_format
        va.tick_labels.number_format_is_linked = False
        if value_axis_max is not None:
            va.maximum_scale = value_axis_max
            va.minimum_scale = 0
        ca = chart.category_axis
        ca.format.line.color.rgb = GRID
        ca.tick_labels.font.size = Pt(12)
        ca.tick_labels.font.color.rgb = INK
        ca.has_major_gridlines = False
    except (ValueError, AttributeError):
        pass


def label_series(plot, number_format='0.0"%"', size=12, color=INK, position=XL_LABEL_POSITION.OUTSIDE_END):
    plot.has_data_labels = True
    dl = plot.data_labels
    dl.number_format = number_format
    dl.number_format_is_linked = False
    dl.font.size = Pt(size)
    dl.font.bold = True
    dl.font.color.rgb = color
    if position is not None:
        dl.position = position


def color_series(chart, colors):
    for series, c in zip(chart.plots[0].series, colors):
        series.format.fill.solid()
        series.format.fill.fore_color.rgb = c
        series.format.line.color.rgb = WHITE
        series.format.line.width = Pt(1.5)


def card(slide, x, y, w, h, fill=ICE):
    return shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h, fill)


# ================================================================ 1. Title
s = prs.slides.add_slide(BLANK)
bg(s, NAVY)
# Big decorative circles (motif) on the right
shape(s, MSO_SHAPE.OVAL, 8.6, -1.2, 6.4, 6.4, NAVY_2)
shape(s, MSO_SHAPE.OVAL, 10.4, 3.9, 3.6, 3.6, SKY)
text(s, 10.4, 5.05, 3.6, 0.6, "43.9%", size=40, color=WHITE, bold=True, align=PP_ALIGN.CENTER, font=HEAD_FONT)
text(s, 10.4, 5.8, 3.6, 0.4, "of passengers satisfied", size=14, color=WHITE, align=PP_ALIGN.CENTER)
text(s, 0.8, 1.3, 8, 0.4, "MID PROJECT  ·  DATA ANALYSIS", size=14, color=LIGHT_MUTED, bold=True)
text(s, 0.8, 1.9, 8.6, 2.0, "Airline Passenger Satisfaction", size=54, color=WHITE, bold=True, font=HEAD_FONT)
text(s, 0.8, 4.0, 8, 0.9, "What makes passengers happy — and what drives them away?", size=22, color=LIGHT_MUTED,
     italic=True)
text(s, 0.8, 5.6, 8, 0.4, "Hisham Mohamed", size=18, color=WHITE, bold=True)
text(s, 0.8, 6.05, 8, 0.4, "Epsilon AI  ·  Data Science Program", size=14, color=LIGHT_MUTED)
s.notes_slide.notes_text_frame.text = (
    "Welcome. This project analyses an airline passenger satisfaction survey of about 26,000 passengers to find "
    "out who is satisfied, who is not, and which parts of the journey make the difference.")

# ================================================================ 2. Agenda
s = content_slide("Agenda", "Roadmap")
agenda = [("The domain & the data", "What the survey measures and what each column means"),
          ("Data problems & cleaning", "What was wrong with the raw data and how it was fixed"),
          ("Research questions", "Six questions that guide the analysis"),
          ("Insights", "Who is satisfied, which services matter, delays and distance"),
          ("Recommendations & conclusion", "What the airline should do next")]
for i, (head, sub) in enumerate(agenda):
    y = 1.75 + i * 1.0
    badge(s, 0.9, y, str(i + 1), fill=SKY if i % 2 == 0 else NAVY)
    text(s, 1.75, y - 0.02, 9, 0.4, head, size=22, color=NAVY, bold=True)
    text(s, 1.75, y + 0.4, 9, 0.35, sub, size=14, color=MUTED)
shape(s, MSO_SHAPE.OVAL, 10.0, 2.0, 2.8, 2.8, ICE)
text(s, 10.0, 2.85, 2.8, 0.6, "6", size=48, color=SKY, bold=True, align=PP_ALIGN.CENTER, font=HEAD_FONT)
text(s, 10.0, 3.6, 2.8, 0.4, "questions answered", size=14, color=NAVY, align=PP_ALIGN.CENTER)
s.notes_slide.notes_text_frame.text = "Quick overview of the five parts of the talk."

# ================================================================ 3. Domain
s = content_slide("The domain: airline customer experience", "The domain & the data")
text(s, 0.6, 1.65, 5.6, 3.5, [
    ("Airlines compete on experience. A dissatisfied passenger is likely to book with a competitor next time.",
     {"space_after": 14}),
    ("This survey records, for each passenger:", {"space_after": 8}),
    ([("Who ", {"bold": True, "color": SKY}), ("they are — gender, age, loyalty", {})], {"space_after": 6}),
    ([("How ", {"bold": True, "color": SKY}), ("they travelled — purpose, class, distance, delays", {})],
     {"space_after": 6}),
    ([("What ", {"bold": True, "color": SKY}), ("they thought — 14 service ratings (1–5)", {})], {"space_after": 6}),
    ([("Overall ", {"bold": True, "color": SKY}), ("satisfaction — the target", {})], {}),
], size=16, color=INK)
stats_cards = [("25,976", "passengers surveyed"), ("25", "columns in the raw file"),
               ("14", "services rated 1–5"), ("2", "target classes")]
for i, (num, lab) in enumerate(stats_cards):
    x = 6.9 + (i % 2) * 3.0
    y = 1.75 + (i // 2) * 2.4
    card(s, x, y, 2.7, 2.1)
    text(s, x, y + 0.35, 2.7, 0.9, num, size=44, color=NAVY, bold=True, align=PP_ALIGN.CENTER, font=HEAD_FONT)
    text(s, x + 0.2, y + 1.35, 2.3, 0.5, lab, size=14, color=MUTED, align=PP_ALIGN.CENTER)
s.notes_slide.notes_text_frame.text = (
    "The dataset is the Airline Passenger Satisfaction survey. Each row is one passenger: their profile, "
    "trip details, 14 service ratings and whether they were satisfied overall.")

# ================================================================ 4. Columns
s = content_slide("What do the columns represent?", "The domain & the data")
groups = [
    ("Passenger profile", SKY, ["id — unique passenger id", "Gender — Female / Male",
                                "Customer Type — loyal / disloyal", "Age — 7 to 85 years"]),
    ("Trip details", NAVY, ["Type of Travel — business / personal", "Class — Business, Eco, Eco Plus",
                            "Flight Distance — miles", "Departure & Arrival Delay — minutes"]),
    ("Service ratings (1–5, 0 = N/A)", SKY, ["Wifi · Time convenience · Online booking", "Gate location · Food & drink",
                                             "Online boarding · Seat comfort", "Entertainment · On-board service",
                                             "Leg room · Baggage · Check-in", "Inflight service · Cleanliness"]),
    ("Target", ORANGE, ["satisfaction —", "satisfied  or", "neutral or dissatisfied"]),
]
widths = [2.95, 3.05, 3.65, 2.1]
x = 0.6
for (head, color, items), w in zip(groups, widths):
    card(s, x, 1.7, w, 4.95)
    shape(s, MSO_SHAPE.OVAL, x + 0.25, 1.95, 0.3, 0.3, color)
    text(s, x + 0.7, 1.93, w - 0.85, 0.7, head, size=16, color=NAVY, bold=True)
    text(s, x + 0.25, 2.85, w - 0.45, 3.7, [(it, {"space_after": 10}) for it in items], size=14, color=INK)
    x += w + 0.22
s.notes_slide.notes_text_frame.text = (
    "Four groups of columns: profile, trip details, the 14 service ratings and the target. Important: a rating of 0 "
    "is not a bad score — it means the question did not apply to the passenger.")

# ================================================================ 5. Problems
s = content_slide("Problems found in the raw data", "Data quality")
zero_total = int((raw.loc[:, "Inflight wifi service":"Cleanliness"] == 0).sum().sum())
problems = [
    (f"{raw['Arrival Delay in Minutes'].isna().sum()}", "missing values", "in Arrival Delay — also stored as float"),
    (f"{zero_total:,}", "'0' ratings", "0 = 'Not Applicable', outside the 1–5 scale; would drag averages down"),
    ("4", "inconsistent columns", "'disloyal Customer', 'Business travel', 'Eco', lower-case target"),
    ("1", "useless column", "'Unnamed: 0' — a leftover row index"),
    ("3,538", "delay outliers", "Up to 1,128 min; verified as genuine (r = 0.96 between delays)"),
    ("14", "messy names", "Spaces, slashes, mixed case — e.g. 'Departure/Arrival time convenient'"),
]
for i, (num, head, desc) in enumerate(problems):
    col, row = i % 3, i // 3
    x, y = 0.6 + col * 4.1, 1.75 + row * 2.55
    card(s, x, y, 3.85, 2.25)
    text(s, x + 0.3, y + 0.25, 3.3, 0.7, num, size=36, color=ORANGE, bold=True, font=HEAD_FONT)
    text(s, x + 0.3, y + 0.98, 3.3, 0.4, head, size=16, color=NAVY, bold=True)
    text(s, x + 0.3, y + 1.38, 3.3, 0.8, desc, size=13, color=MUTED)
s.notes_slide.notes_text_frame.text = (
    "The data is not clean. The most important hidden problem is the 4,202 zero ratings: they look like the worst "
    "score but actually mean not applicable. There are also 83 missing arrival delays, inconsistent labels, a useless "
    "index column and extreme delay values, which I checked and found to be genuine.")

# ================================================================ 6. Cleaning
s = content_slide("How the data was cleaned", "Data cleaning")
rows = [("Missing arrival delay (83)", "Filled with the same passenger's departure delay (correlation 0.96)"),
        ("Rating 0 = Not Applicable", "Replaced with NaN → ignored in averages; counted per passenger"),
        ("Inconsistent labels", "Standardised: Disloyal Customer, Business Travel, Economy, Satisfied"),
        ("Float delay column", "Converted to integer minutes"),
        ("Useless index & messy names", "Dropped 'Unnamed: 0'; renamed all columns to snake_case"),
        ("Extreme delays / distances", "Kept (real events) and grouped into bins"),
        ("No segment columns", "Added age, distance & delay groups, average rating, satisfied flag")]
tbl = s.shapes.add_table(len(rows) + 1, 2, Inches(0.6), Inches(1.7), Inches(12.1), Inches(4.9)).table
tbl.columns[0].width = Inches(4.0)
tbl.columns[1].width = Inches(8.1)
for c, head in enumerate(["Problem", "Action taken"]):
    cell = tbl.cell(0, c)
    cell.fill.solid()
    cell.fill.fore_color.rgb = NAVY
    cell.text = head
    para = cell.text_frame.paragraphs[0]
    para.runs[0].font.size, para.runs[0].font.bold = Pt(15), True
    para.runs[0].font.color.rgb, para.runs[0].font.name = WHITE, FONT
for r, (prob, act) in enumerate(rows, start=1):
    for c, val in enumerate([prob, act]):
        cell = tbl.cell(r, c)
        cell.fill.solid()
        cell.fill.fore_color.rgb = ICE if r % 2 else WHITE
        cell.text = val
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        run = cell.text_frame.paragraphs[0].runs[0]
        run.font.size, run.font.name = Pt(14), FONT
        run.font.bold = c == 0
        run.font.color.rgb = NAVY if c == 0 else INK
s.notes_slide.notes_text_frame.text = (
    "Each problem was handled with Define, Code, Test. Instead of the median, missing arrival delays were filled with "
    "the passenger's own departure delay because the two are almost perfectly correlated. Nothing was deleted: "
    "the final data has 25,976 rows and 30 columns.")

# ================================================================ 7. Questions
s = content_slide("Six questions guide the analysis", "Research questions")
questions = ["What share of passengers are satisfied?",
             "Which passenger segments are more or less satisfied?",
             "Which services matter most for satisfaction?",
             "Do departure and arrival delays reduce satisfaction?",
             "Does flight distance matter — and does it depend on class?",
             "Which travel type × class combination is happiest / unhappiest?"]
for i, q in enumerate(questions):
    col, row = i % 2, i // 2
    x, y = 0.6 + col * 6.15, 1.8 + row * 1.65
    card(s, x, y, 5.9, 1.35)
    badge(s, x + 0.3, y + 0.38, f"Q{i + 1}", fill=SKY if i % 2 == 0 else NAVY, size=0.62, font_size=15)
    text(s, x + 1.15, y + 0.2, 4.55, 0.95, q, size=17, color=INK, anchor=MSO_ANCHOR.MIDDLE)
s.notes_slide.notes_text_frame.text = "These six questions structure the analysis in the notebook and in the dashboard."

# ================================================================ 8. Section divider: insights
s = dark_slide()
shape(s, MSO_SHAPE.OVAL, 9.3, 1.4, 4.6, 4.6, NAVY_2)
text(s, 0.8, 2.5, 9, 0.5, "PART 2", size=16, color=LIGHT_MUTED, bold=True)
text(s, 0.8, 3.0, 9, 1.2, "Insights from the data", size=48, color=WHITE, bold=True, font=HEAD_FONT)
text(s, 0.8, 4.2, 8.5, 0.6, "Each finding is backed by a statistical test (chi-square, Mann-Whitney, correlation)",
     size=18, color=LIGHT_MUTED, italic=True)

# ================================================================ 9. Q1
s = content_slide("Most passengers are not satisfied", "Q1 · Overall satisfaction")
counts = df["satisfaction"].value_counts()
cd = CategoryChartData()
cd.categories = ["Satisfied", "Neutral or Dissatisfied"]
cd.add_series("Passengers", [int(counts["Satisfied"]), int(counts["Neutral or Dissatisfied"])])
gf = s.shapes.add_chart(XL_CHART_TYPE.DOUGHNUT, Inches(0.6), Inches(1.6), Inches(6.0), Inches(5.2), cd)
ch = gf.chart
ch.has_title = False
ch.has_legend = True
ch.legend.position = XL_LEGEND_POSITION.BOTTOM
ch.legend.include_in_layout = False
ch.legend.font.size = Pt(14)
ch.font.name = FONT
pts = ch.plots[0].series[0].points
for idx, colr in enumerate([SKY, ORANGE]):
    pts[idx].format.fill.solid()
    pts[idx].format.fill.fore_color.rgb = colr
    pts[idx].format.line.color.rgb = WHITE
ch.plots[0].has_data_labels = True
ch.plots[0].data_labels.show_percentage = True
ch.plots[0].data_labels.show_value = False
ch.plots[0].data_labels.number_format = "0.0%"
ch.plots[0].data_labels.number_format_is_linked = False
ch.plots[0].data_labels.font.size = Pt(16)
ch.plots[0].data_labels.font.bold = True
ch.plots[0].data_labels.font.color.rgb = WHITE
text(s, 7.3, 2.0, 5.4, 1.2, "43.9%", size=72, color=SKY, bold=True, font=HEAD_FONT)
text(s, 7.3, 3.25, 5.4, 0.5, "of passengers are satisfied", size=20, color=NAVY, bold=True)
text(s, 7.3, 4.0, 5.4, 2.0, [
    ("14,573 passengers (56.1 %) were neutral or dissatisfied.", {"space_after": 10}),
    ("The two groups are reasonably balanced, so comparisons between them are reliable.", {}),
], size=16, color=INK)
s.notes_slide.notes_text_frame.text = "Answer to Q1: only 43.9 percent are satisfied. The majority are not happy."

# ================================================================ 10. Q2 segments
s = content_slide("Class and travel purpose split passengers sharply", "Q2 · Passenger segments")
segs = [("Business class", rate("class")["Business"]), ("Economy Plus", rate("class")["Economy Plus"]),
        ("Economy", rate("class")["Economy"]), ("Business travel", rate("type_of_travel")["Business Travel"]),
        ("Personal travel", rate("type_of_travel")["Personal Travel"]),
        ("Loyal customer", rate("customer_type")["Loyal Customer"]),
        ("Disloyal customer", rate("customer_type")["Disloyal Customer"]),
        ("Female", rate("gender")["Female"]), ("Male", rate("gender")["Male"])]
cd = CategoryChartData()
cd.categories = [n for n, _ in segs]
cd.add_series("Satisfied (%)", [round(v, 1) for _, v in segs])
gf = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(0.6), Inches(1.6), Inches(8.4), Inches(5.2), cd)
ch = gf.chart
style_chart(ch, value_axis_max=100)
ch.plots[0].gap_width = 45
label_series(ch.plots[0])
color_series(ch, [SKY])
for i, (_, v) in enumerate(segs):
    p = ch.plots[0].series[0].points[i]
    p.format.fill.solid()
    p.format.fill.fore_color.rgb = SKY if v >= 43.9 else ORANGE
ch.category_axis.tick_labels.font.size = Pt(11)
card(s, 9.4, 1.75, 3.35, 4.9)
text(s, 9.65, 1.95, 2.9, 4.6, [
    ("Strongest drivers", {"bold": True, "color": NAVY, "size": 17, "space_after": 8}),
    ([("Class ", {"bold": True}), ("— Business 69.5 % vs Economy 19.4 %", {})], {"space_after": 8}),
    ([("Travel type ", {"bold": True}), ("— business 58.8 % vs personal 10.0 %", {})], {"space_after": 8}),
    ([("Loyalty ", {"bold": True}), ("— loyal 48.1 % vs disloyal 25.2 %", {})], {"space_after": 8}),
    ([("Gender ", {"bold": True}), ("— no difference (p = 0.24)", {})], {"space_after": 8}),
    ("Blue = above the 43.9 % average", {"size": 12, "color": MUTED, "italic": True}),
], size=14, color=INK)
s.notes_slide.notes_text_frame.text = (
    "Satisfaction rate per segment. Class and travel type have the strongest association (Cramer's V 0.50 and 0.45). "
    "Gender has no significant effect.")

# ================================================================ 11. Age
s = content_slide("Middle-aged passengers are the most satisfied", "Q2 · Age")
age_order = ["Child (<18)", "Young Adult (18-29)", "Adult (30-44)", "Middle-Aged (45-59)", "Senior (60+)"]
ar = rate("age_group", age_order)
cd = CategoryChartData()
cd.categories = ["Child\n<18", "Young adult\n18-29", "Adult\n30-44", "Middle-aged\n45-59", "Senior\n60+"]
cd.add_series("Satisfied (%)", [round(v, 1) for v in ar.values])
gf = s.shapes.add_chart(XL_CHART_TYPE.LINE_MARKERS, Inches(0.6), Inches(1.6), Inches(8.4), Inches(5.2), cd)
ch = gf.chart
style_chart(ch, value_axis_max=80)
ser = ch.plots[0].series[0]
ser.format.line.color.rgb = SKY
ser.format.line.width = Pt(3.5)
ser.smooth = False
ser.marker.size = 12
ser.marker.format.fill.solid()
ser.marker.format.fill.fore_color.rgb = SKY
ser.marker.format.line.color.rgb = WHITE
label_series(ch.plots[0], position=XL_LABEL_POSITION.ABOVE)
card(s, 9.4, 1.75, 3.35, 4.9)
text(s, 9.65, 1.95, 2.9, 4.6, [
    ("57.7 %", {"bold": True, "color": SKY, "size": 40, "font": HEAD_FONT}),
    ("of 45–59-year-olds are satisfied", {"bold": True, "color": NAVY, "space_after": 14}),
    ("Children (18 %), young adults (36 %) and seniors (27 %) are the least satisfied — groups that mostly travel "
     "for personal reasons in Economy.", {"space_after": 10}),
    ("Median age: satisfied 43 vs dissatisfied 37 (Mann-Whitney p < 0.001)", {"size": 12, "color": MUTED}),
], size=14, color=INK)
s.notes_slide.notes_text_frame.text = "Satisfaction rises with age until about 60, then drops again."

# ================================================================ 12. Q3 services
s = content_slide("Boarding, entertainment and wifi make the difference", "Q3 · Services")
means = df.groupby("satisfaction")[SERVICES].mean().T
means["gap"] = means["Satisfied"] - means["Neutral or Dissatisfied"]
means = means.sort_values("gap")
nice = {"departure_arrival_time_convenient": "Time convenience", "inflight_wifi_service": "Inflight wifi",
        "ease_of_online_booking": "Online booking", "on_board_service": "On-board service",
        "leg_room_service": "Leg room", "checkin_service": "Check-in"}
labels = [nice.get(c, c.replace("_", " ").capitalize()) for c in means.index]
cd = CategoryChartData()
cd.categories = labels
cd.add_series("Dissatisfied", [round(v, 2) for v in means["Neutral or Dissatisfied"]])
cd.add_series("Satisfied", [round(v, 2) for v in means["Satisfied"]])
gf = s.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED, Inches(0.6), Inches(1.5), Inches(8.0), Inches(5.4), cd)
ch = gf.chart
style_chart(ch, legend=True, value_axis_max=5, number_format="0", font_size=11)
ch.category_axis.tick_labels.font.size = Pt(11)
ch.value_axis.major_unit = 1
ch.plots[0].gap_width = 40
ch.plots[0].overlap = 0
color_series(ch, [ORANGE, SKY])
top = means.sort_values("gap", ascending=False)
card(s, 9.0, 1.75, 3.75, 4.9)
text(s, 9.25, 1.95, 3.3, 4.6, [
    ("Biggest rating gaps", {"bold": True, "color": NAVY, "size": 17, "space_after": 8}),
    ([("+1.44 ", {"bold": True, "color": SKY}), ("Online boarding", {})], {"space_after": 4}),
    ([("+1.08 ", {"bold": True, "color": SKY}), ("Inflight entertainment", {})], {"space_after": 4}),
    ([("+1.00 ", {"bold": True, "color": SKY}), ("Inflight wifi", {})], {"space_after": 4}),
    ([("+0.92 ", {"bold": True, "color": SKY}), ("Seat comfort", {})], {"space_after": 12}),
    ([("Wifi is the lowest-rated service (2.81) ", {"bold": True, "color": ORANGE}),
      ("— yet 99 % of passengers who rate it 5 are satisfied.", {})], {"space_after": 10}),
    ("Gate location and time convenience show no link.", {"size": 13, "color": MUTED}),
], size=14, color=INK)
s.notes_slide.notes_text_frame.text = (
    "Average rating of each service for satisfied vs dissatisfied passengers. The largest gaps show which services "
    "separate happy from unhappy passengers. Online boarding has the strongest correlation, 0.59.")

# ================================================================ 13. Q4 delays
s = content_slide("Delays lower satisfaction — but only modestly", "Q4 · Delays")
dr = rate("delay_group", ["On time", "1-15 min", "16-60 min", "> 60 min"])
cd = CategoryChartData()
cd.categories = ["On time", "1–15 min", "16–60 min", "> 60 min"]
cd.add_series("Satisfied (%)", [round(v, 1) for v in dr.values])
gf = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(0.6), Inches(1.6), Inches(7.6), Inches(5.2), cd)
ch = gf.chart
style_chart(ch, value_axis_max=60)
ch.plots[0].gap_width = 60
label_series(ch.plots[0])
color_series(ch, [SKY])
ch.category_axis.has_title = True
ch.category_axis.axis_title.text_frame.text = "Arrival delay"
ch.category_axis.axis_title.text_frame.paragraphs[0].runs[0].font.size = Pt(12)
stat_cards = [("−12.7 pts", "from on-time to > 1 hour late"), ("56.3 %", "of flights arrive on time"),
              ("V = 0.10", "very weak association (p < 0.001)")]
for i, (num, lab) in enumerate(stat_cards):
    y = 1.75 + i * 1.65
    card(s, 8.7, y, 4.05, 1.4)
    text(s, 8.95, y + 0.18, 3.6, 0.6, num, size=28, color=NAVY if i else ORANGE, bold=True, font=HEAD_FONT)
    text(s, 8.95, y + 0.82, 3.6, 0.45, lab, size=14, color=MUTED)
s.notes_slide.notes_text_frame.text = (
    "Satisfaction drops steadily from 47.9 percent on time to 35.2 percent for delays over an hour. Significant, but "
    "much weaker than class or service quality.")

# ================================================================ 14. Q5 distance x class
s = content_slide("Distance matters only through class", "Q5 · Flight distance")
dist_order = ["Short-haul (<800 mi)", "Medium-haul (800-2000 mi)", "Long-haul (>2000 mi)"]
pv = df.pivot_table(index="distance_group", columns="class", values="is_satisfied", aggfunc="mean") * 100
pv = pv.reindex(dist_order)
cd = CategoryChartData()
cd.categories = ["Short-haul\n< 800 mi", "Medium-haul\n800–2,000 mi", "Long-haul\n> 2,000 mi"]
for cls in ["Economy", "Economy Plus", "Business"]:
    cd.add_series(cls, [round(v, 1) for v in pv[cls].values])
gf = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(0.6), Inches(1.5), Inches(8.0), Inches(5.3), cd)
ch = gf.chart
style_chart(ch, legend=True, value_axis_max=100)
ch.plots[0].gap_width = 60
ch.plots[0].overlap = -5
label_series(ch.plots[0], number_format='0"%"', size=11)
color_series(ch, [ORANGE, RGBColor(0xED, 0xA1, 0x00), SKY])
card(s, 9.0, 1.75, 3.75, 4.9)
text(s, 9.25, 1.95, 3.3, 4.6, [
    ("A hidden confounder", {"bold": True, "color": NAVY, "size": 17, "space_after": 8}),
    ("Overall, long-haul passengers look happier (70 % vs 34 % short-haul).", {"space_after": 10}),
    ("But long-haul flights are mostly flown in Business class (median 1,590 mi vs ~600 mi in Economy).",
     {"space_after": 10}),
    ([("Within Economy, distance changes almost nothing (17–21 %). ", {"bold": True}),
      ("Only Business class improves on long flights (60 → 77 %).", {})], {}),
], size=14, color=INK)
s.notes_slide.notes_text_frame.text = (
    "This is a good example of why multivariate analysis matters: the distance effect disappears when we control "
    "for class.")

# ================================================================ 15. Q6 combination
s = content_slide("Leisure travellers are unhappy in every class", "Q6 · Travel type × class")
tc = df.pivot_table(index="type_of_travel", columns="class", values="is_satisfied", aggfunc="mean") * 100
tn = df.pivot_table(index="type_of_travel", columns="class", values="id", aggfunc="count")
classes = ["Economy", "Economy Plus", "Business"]
travel = ["Business Travel", "Personal Travel"]
x0, y0, cw, chh = 3.1, 2.15, 2.15, 1.75
for j, cls in enumerate(classes):
    text(s, x0 + j * (cw + 0.12), y0 - 0.5, cw, 0.4, cls, size=16, color=NAVY, bold=True, align=PP_ALIGN.CENTER)
for i, tr in enumerate(travel):
    text(s, 0.6, y0 + i * (chh + 0.12) + 0.55, 2.4, 0.6, tr, size=16, color=NAVY, bold=True)
    for j, cls in enumerate(classes):
        v = tc.loc[tr, cls]
        shade = SKY if v > 60 else (RGBColor(0x9C, 0xC3, 0xEE) if v > 25 else ICE)
        tx = WHITE if v > 60 else NAVY
        x, y = x0 + j * (cw + 0.12), y0 + i * (chh + 0.12)
        card(s, x, y, cw, chh, fill=shade)
        text(s, x, y + 0.35, cw, 0.7, f"{v:.1f}%", size=32, color=tx, bold=True, align=PP_ALIGN.CENTER,
             font=HEAD_FONT)
        text(s, x, y + 1.1, cw, 0.4, f"n = {tn.loc[tr, cls]:,}", size=12, color=tx, align=PP_ALIGN.CENTER)
text(s, 0.6, 6.0, 8.6, 0.6, "Share of satisfied passengers per combination (darker = more satisfied)", size=12,
     color=MUTED, italic=True)
card(s, 10.0, 1.65, 2.75, 4.7)
text(s, 10.2, 1.85, 2.4, 4.4, [
    ("Happiest", {"bold": True, "color": SKY, "size": 17}),
    ("Business travellers in Business class — 72 %", {"space_after": 14}),
    ("Unhappiest", {"bold": True, "color": ORANGE, "size": 17}),
    ("Personal travellers — about 9–10 % in every class, even Business", {"space_after": 14}),
    ("Loyalty doesn't help them: loyal leisure travellers are only 10 % satisfied.", {"size": 13, "color": MUTED}),
], size=14, color=INK)
s.notes_slide.notes_text_frame.text = (
    "The core business segment is well served. Leisure travellers are the big problem, regardless of class.")

# ================================================================ 16. Drivers
s = content_slide("What drives satisfaction the most?", "Summary of drivers")
drivers = [("Class", cramers_v("class")), ("Type of travel", cramers_v("type_of_travel")),
           ("Age group", cramers_v("age_group")), ("Customer type", cramers_v("customer_type")),
           ("Arrival delay", cramers_v("delay_group")), ("Gender", cramers_v("gender"))][::-1]
cd = CategoryChartData()
cd.categories = [d for d, _ in drivers]
cd.add_series("Cramér's V", [round(v, 2) for _, v in drivers])
gf = s.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED, Inches(0.6), Inches(1.6), Inches(7.6), Inches(5.0), cd)
ch = gf.chart
style_chart(ch, value_axis_max=0.6, number_format="0.0")
ch.plots[0].gap_width = 50
label_series(ch.plots[0], number_format="0.00")
color_series(ch, [SKY])
ch.category_axis.tick_labels.font.size = Pt(13)
card(s, 8.7, 1.75, 4.05, 4.9)
text(s, 8.95, 1.95, 3.6, 4.6, [
    ("Reading the chart", {"bold": True, "color": NAVY, "size": 17, "space_after": 8}),
    ("Cramér's V measures how strongly each factor is linked with satisfaction (0 = none, 0.5 = strong).",
     {"space_after": 12}),
    ([("Service quality ", {"bold": True}), ("(online boarding r = 0.59) is as strong as class.", {})],
     {"space_after": 12}),
    ([("Delays and gender ", {"bold": True}), ("matter little.", {})], {}),
], size=14, color=INK)
s.notes_slide.notes_text_frame.text = "Ranking of all the factors by strength of association with satisfaction."

# ================================================================ 17. Recommendations
s = content_slide("Recommendations for the airline", "What to do next")
recs = [("Fix the digital journey", "Wifi, online boarding and booking: lowest rated, yet top drivers"),
        ("Upgrade Economy", "Fewer than 1 in 5 Economy passengers are satisfied — seats, entertainment, leg room"),
        ("Win over leisure travellers", "Only ~10 % satisfied: family services, flexible fares, entertainment"),
        ("Re-engage disloyal customers", "25 % satisfied — targeted loyalty offers"),
        ("Keep delays short", "Smaller effect, but satisfaction falls with every delay band")]
for i, (head, desc) in enumerate(recs):
    y = 1.7 + i * 1.03
    card(s, 0.6, y, 12.15, 0.88)
    badge(s, 0.85, y + 0.16, str(i + 1), fill=SKY if i < 3 else NAVY)
    text(s, 1.65, y + 0.12, 3.6, 0.65, head, size=18, color=NAVY, bold=True, anchor=MSO_ANCHOR.MIDDLE)
    text(s, 5.3, y + 0.12, 7.2, 0.65, desc, size=15, color=INK, anchor=MSO_ANCHOR.MIDDLE)
s.notes_slide.notes_text_frame.text = "Five actions, ordered by expected impact."

# ================================================================ 18. Conclusion
s = dark_slide()
text(s, 0.8, 0.7, 11, 0.5, "CONCLUSION", size=16, color=LIGHT_MUTED, bold=True)
text(s, 0.8, 1.2, 11.8, 1.0, "Satisfaction is earned in the cabin and online — not by distance or gender",
     size=34, color=WHITE, bold=True, font=HEAD_FONT)
concl = [("43.9 %", "satisfied overall — room to grow"),
         ("72 % vs 10 %", "business travellers in Business class vs leisure travellers"),
         ("+1.44", "online boarding: the single biggest service gap")]
for i, (num, lab) in enumerate(concl):
    x = 0.8 + i * 4.0
    shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, x, 3.0, 3.7, 2.4, NAVY_2)
    text(s, x + 0.3, 3.3, 3.1, 0.9, num, size=36, color=WHITE, bold=True, font=HEAD_FONT)
    text(s, x + 0.3, 4.25, 3.1, 1.0, lab, size=15, color=LIGHT_MUTED)
text(s, 0.8, 5.9, 11.8, 0.8,
     "The data was cleaned (N/A ratings, missing delays, inconsistent labels), explored across 10+ variables with "
     "6 kinds of charts, and every insight was verified with a statistical test.", size=15, color=LIGHT_MUTED)
s.notes_slide.notes_text_frame.text = "Wrap up: the three numbers to remember."

# ================================================================ 19. Thank you
s = dark_slide()
shape(s, MSO_SHAPE.OVAL, 8.9, 1.2, 5.0, 5.0, NAVY_2)
shape(s, MSO_SHAPE.OVAL, 10.6, 4.4, 2.2, 2.2, SKY)
text(s, 0.8, 2.3, 8, 1.2, "Thank you", size=60, color=WHITE, bold=True, font=HEAD_FONT)
text(s, 0.8, 3.6, 8, 0.6, "Questions & discussion", size=24, color=LIGHT_MUTED, italic=True)
text(s, 0.8, 4.8, 8, 1.2, [
    ("GitHub: github.com/hahmed1936/Mid-Project", {"space_after": 6}),
    ("Interactive dashboard: Streamlit (link in the repository README)", {}),
], size=16, color=WHITE)

prs.save(OUT)
print("saved", OUT, len(prs.slides), "slides")
