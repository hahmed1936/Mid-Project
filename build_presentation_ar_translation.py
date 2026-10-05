"""Arabic translation of the discussion deck (same 19 slides, numbers and charts as the English version)."""
from pathlib import Path

import numpy as np
import pandas as pd
from scipy import stats
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_LABEL_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt

from pptx_ar_utils import (ArabicDeck, style_chart, label_series, color_series, NAVY, NAVY_2, SKY, ORANGE, YELLOW,
                           ICE, WHITE, INK, MUTED, LIGHT_MUTED, FONT)

PROJECT = Path(__file__).parent
OUT = PROJECT / "Airline_Satisfaction_Presentation_AR.pptx"

# ---------------------------------------------------------------- data (same calculations as the English deck)
df = pd.read_csv(PROJECT / "Data_Cleaned.csv")
raw = pd.read_csv(PROJECT / "Data.csv")
SERVICES = list(df.loc[:, "inflight_wifi_service":"cleanliness"].columns)


def rate(col, order=None):
    s = df.groupby(col)["is_satisfied"].mean() * 100
    return s.reindex(order) if order else s


def cramers_v(col):
    t = pd.crosstab(df[col], df["satisfaction"])
    chi2 = stats.chi2_contingency(t)[0]
    return np.sqrt(chi2 / (t.values.sum() * (min(t.shape) - 1)))


deck = ArabicDeck("رضا ركاب الطيران  ·  المشروع المرحلي")
T = deck.text

# ================================================================ 1. Title
s = deck.dark_slide()
deck.shape(s, MSO_SHAPE.OVAL, 8.6, -1.2, 6.4, 6.4, NAVY_2)
deck.shape(s, MSO_SHAPE.OVAL, 10.4, 3.9, 3.6, 3.6, SKY)
T(s, 10.4, 5.05, 3.6, 0.6, "43.9%", size=40, color=WHITE, bold=True, rtl=False, align=PP_ALIGN.CENTER)
T(s, 10.4, 5.8, 3.6, 0.4, "من الركاب راضون", size=15, color=WHITE, align=PP_ALIGN.CENTER)
T(s, 0.8, 1.3, 8.6, 0.4, "المشروع المرحلي  ·  تحليل البيانات", size=15, color=LIGHT_MUTED, bold=True)
T(s, 0.8, 1.9, 8.6, 2.0, "رضا ركاب الطيران", size=56, color=WHITE, bold=True)
T(s, 0.8, 3.6, 8.6, 0.9, "ما الذي يُسعد الركاب — وما الذي يُبعدهم؟", size=24, color=LIGHT_MUTED)
T(s, 0.8, 5.6, 8.6, 0.4, "هشام محمد", size=20, color=WHITE, bold=True)
T(s, 0.8, 6.05, 8.6, 0.4, "Epsilon AI  ·  برنامج علوم البيانات", size=15, color=LIGHT_MUTED)
deck.notes(s, "أهلًا بكم. هذا المشروع يحلل استبيان رضا ركاب طيران لحوالي 26 ألف راكب، لمعرفة من هو الراضي ومن "
              "غير الراضي، وما هي أجزاء الرحلة التي تصنع الفرق.")

# ================================================================ 2. Agenda
s = deck.content_slide("جدول المحتويات", "خريطة العرض")
agenda = [("المجال والبيانات", "ماذا يقيس الاستبيان وماذا يعني كل عمود"),
          ("مشاكل البيانات وتنظيفها", "ما الخطأ في البيانات الخام وكيف تم إصلاحه"),
          ("الأسئلة البحثية", "ستة أسئلة توجّه التحليل"),
          ("النتائج", "من الراضي، أي الخدمات تؤثر، التأخير والمسافة"),
          ("التوصيات والخلاصة", "ما الذي يجب أن تفعله شركة الطيران")]
for i, (head, sub) in enumerate(agenda):
    y = 1.75 + i * 1.0
    deck.badge(s, 0.9, y, str(i + 1), fill=SKY if i % 2 == 0 else NAVY)
    T(s, 1.75, y - 0.02, 7.6, 0.4, head, size=22, color=NAVY, bold=True)
    T(s, 1.75, y + 0.42, 7.6, 0.35, sub, size=14, color=MUTED)
deck.shape(s, MSO_SHAPE.OVAL, 10.0, 2.0, 2.8, 2.8, ICE)
T(s, 10.0, 2.85, 2.8, 0.6, "6", size=48, color=SKY, bold=True, rtl=False, align=PP_ALIGN.CENTER)
T(s, 10.0, 3.6, 2.8, 0.4, "أسئلة تمت الإجابة عليها", size=14, color=NAVY, align=PP_ALIGN.CENTER)
deck.notes(s, "نظرة سريعة على الأجزاء الخمسة للعرض.")

# ================================================================ 3. Domain
s = deck.content_slide("المجال: تجربة عملاء شركات الطيران", "المجال والبيانات")
T(s, 0.6, 1.65, 5.6, 3.8, [
    ("شركات الطيران تتنافس على تجربة العميل. الراكب غير الراضي غالبًا سيحجز مع منافس في المرة القادمة.",
     {"space_after": 14}),
    ("يسجّل هذا الاستبيان لكل راكب:", {"space_after": 8}),
    ([("مَن ", {"bold": True, "color": SKY}), ("هو — النوع، العمر، الولاء", {})], {"space_after": 6}),
    ([("كيف ", {"bold": True, "color": SKY}), ("سافر — الغرض، الدرجة، المسافة، التأخير", {})], {"space_after": 6}),
    ([("ماذا ", {"bold": True, "color": SKY}), ("رأى — تقييم 14 خدمة من 1 إلى 5", {})], {"space_after": 6}),
    ([("الرضا العام ", {"bold": True, "color": SKY}), ("— المتغير المستهدف", {})], {}),
], size=16)
for i, (num, lab) in enumerate([("25,976", "راكب شملهم الاستبيان"), ("25", "عمودًا في الملف الخام"),
                                ("14", "خدمة مقيّمة من 1 إلى 5"), ("2", "فئتان للمتغير المستهدف")]):
    x, y = 6.9 + (i % 2) * 3.0, 1.75 + (i // 2) * 2.4
    deck.card(s, x, y, 2.7, 2.1)
    T(s, x, y + 0.35, 2.7, 0.9, num, size=44, color=NAVY, bold=True, rtl=False, align=PP_ALIGN.CENTER)
    T(s, x + 0.2, y + 1.35, 2.3, 0.5, lab, size=14, color=MUTED, align=PP_ALIGN.CENTER)
deck.notes(s, "البيانات هي استبيان رضا ركاب الطيران. كل صف يمثل راكبًا واحدًا: بياناته الشخصية، وتفاصيل رحلته، "
              "وتقييمه لـ 14 خدمة، وهل كان راضيًا بشكل عام أم لا.")

# ================================================================ 4. Columns
s = deck.content_slide("ماذا تمثل الأعمدة؟", "المجال والبيانات")
groups = [
    ("بيانات الراكب", SKY, ["id — رقم تعريف الراكب", "Gender — أنثى / ذكر", "Customer Type — عميل دائم / غير دائم",
                            "Age — من 7 إلى 85 سنة"]),
    ("تفاصيل الرحلة", NAVY, ["Type of Travel — عمل / شخصي", "Class — Business, Eco, Eco Plus",
                             "Flight Distance — بالميل", "Departure / Arrival Delay — بالدقائق"]),
    ("تقييم الخدمات (1–5، و0 = لا ينطبق)", SKY, ["الواي فاي · ملاءمة المواعيد · الحجز أونلاين",
                                                 "موقع البوابة · الطعام والشراب", "الصعود أونلاين · راحة المقعد",
                                                 "الترفيه · الخدمة على متن الطائرة",
                                                 "مساحة الأرجل · الأمتعة · تسجيل الوصول",
                                                 "الخدمة أثناء الرحلة · النظافة"]),
    ("المتغير المستهدف", ORANGE, ["satisfaction —", "راضٍ  أو", "محايد أو غير راضٍ"]),
]
x = 0.6
for (head, color, items), w in zip(groups, [2.95, 3.05, 3.65, 2.1]):
    deck.card(s, x, 1.7, w, 4.95)
    deck.shape(s, MSO_SHAPE.OVAL, x + 0.25, 1.95, 0.3, 0.3, color)
    T(s, x + 0.7, 1.93, w - 0.9, 0.75, head, size=16, color=NAVY, bold=True)
    T(s, x + 0.25, 2.85, w - 0.45, 3.7, [(it, {"space_after": 10}) for it in items], size=14)
    x += w + 0.22
deck.notes(s, "أربع مجموعات من الأعمدة: بيانات الراكب، تفاصيل الرحلة، تقييمات الخدمات الأربع عشرة، والمتغير "
              "المستهدف. ملاحظة مهمة: التقييم صفر ليس تقييمًا سيئًا، بل يعني أن السؤال لا ينطبق على الراكب.")

# ================================================================ 5. Problems
s = deck.content_slide("المشاكل الموجودة في البيانات الخام", "جودة البيانات")
zero_total = int((raw.loc[:, "Inflight wifi service":"Cleanliness"] == 0).sum().sum())
problems = [(f"{raw['Arrival Delay in Minutes'].isna().sum()}", "قيمة ناقصة",
             "في عمود تأخير الوصول — ومخزّن كرقم عشري"),
            (f"{zero_total:,}", "تقييم بصفر", "الصفر يعني «لا ينطبق» وهو خارج مقياس 1–5، ويخفض المتوسطات"),
            ("4", "أعمدة غير متسقة", "مثل disloyal Customer و Business travel و Eco"),
            ("1", "عمود بلا فائدة", "Unnamed: 0 — رقم صف متبقٍ من تصدير سابق"),
            ("3,538", "قيمة متطرفة في التأخير", "حتى 1,128 دقيقة، وتم التأكد أنها حقيقية"),
            ("14", "اسم عمود غير مرتب", "مسافات وشرطات وحروف كبيرة وصغيرة")]
for i, (num, head, desc) in enumerate(problems):
    x, y = 0.6 + (i % 3) * 4.1, 1.75 + (i // 3) * 2.55
    deck.card(s, x, y, 3.85, 2.25)
    T(s, x + 0.3, y + 0.25, 3.3, 0.7, num, size=36, color=ORANGE, bold=True, rtl=False, align=PP_ALIGN.RIGHT)
    T(s, x + 0.3, y + 0.98, 3.3, 0.4, head, size=16, color=NAVY, bold=True)
    T(s, x + 0.3, y + 1.4, 3.3, 0.8, desc, size=13, color=MUTED)
deck.notes(s, "البيانات ليست نظيفة. أهم مشكلة مخفية هي 4,202 تقييم بصفر: تبدو كأسوأ تقييم لكنها تعني «لا ينطبق». "
              "وهناك أيضًا 83 قيمة ناقصة في تأخير الوصول، وتسميات غير متسقة، وعمود فهرس بلا فائدة، وقيم تأخير متطرفة "
              "تحققت منها ووجدتها حقيقية.")

# ================================================================ 6. Cleaning
s = deck.content_slide("كيف تم تنظيف البيانات", "تنظيف البيانات")
rows = [("تأخير وصول ناقص (83)", "مُلئ بتأخير الإقلاع لنفس الراكب (الارتباط 0.96)"),
        ("التقييم 0 = لا ينطبق", "استُبدل بقيمة فارغة NaN فيُتجاهل في المتوسطات، ويُحسب لكل راكب"),
        ("تسميات غير متسقة", "توحيد: Disloyal Customer و Business Travel و Economy و Satisfied"),
        ("عمود تأخير عشري", "تحويله إلى عدد صحيح من الدقائق"),
        ("عمود فهرس وأسماء غير مرتبة", "حذف Unnamed: 0 وتحويل كل الأسماء إلى snake_case"),
        ("تأخيرات ومسافات متطرفة", "الإبقاء عليها (أحداث حقيقية) وتقسيمها إلى فئات"),
        ("لا توجد أعمدة للشرائح", "إضافة فئات العمر والمسافة والتأخير، ومتوسط التقييم، وعلامة الرضا")]
tbl = deck.table(s, len(rows) + 1, 2, 0.6, 1.7, 12.1, 4.9)
header = ["المشكلة", "الإجراء المتخذ"]
for r, values in enumerate([header] + rows):
    for c, val in enumerate(values[::-1] if deck.mirror else values):   # problem column on the right
        cell = tbl.cell(r, c)
        cell.fill.solid()
        cell.fill.fore_color.rgb = NAVY if r == 0 else (ICE if r % 2 else WHITE)
        cell.text = val
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        para = cell.text_frame.paragraphs[0]
        para.alignment = PP_ALIGN.RIGHT
        para._p.get_or_add_pPr().set("rtl", "1")
        run = para.runs[0]
        run.font.size, run.font.name = Pt(15 if r == 0 else 14), FONT
        is_problem_col = (c == 1) if deck.mirror else (c == 0)
        run.font.bold = r == 0 or is_problem_col
        run.font.color.rgb = WHITE if r == 0 else (NAVY if is_problem_col else INK)
action_width, problem_width = Inches(8.1), Inches(4.0)
tbl.columns[0].width, tbl.columns[1].width = ((action_width, problem_width) if deck.mirror
                                            else (problem_width, action_width))
deck.notes(s, "تم التعامل مع كل مشكلة بطريقة: تعريف ثم كود ثم اختبار. بدل الوسيط، تم ملء تأخير الوصول الناقص بتأخير "
              "الإقلاع لنفس الراكب لأنهما مرتبطان بشكل شبه تام. لم يُحذف أي صف: البيانات النهائية فيها 25,976 صفًا و30 عمودًا.")

# ================================================================ 7. Questions
s = deck.content_slide("ستة أسئلة توجّه التحليل", "الأسئلة البحثية")
questions = ["ما نسبة الركاب الراضين؟", "أي شرائح الركاب أكثر أو أقل رضا؟", "أي الخدمات أكثر تأثيرًا على الرضا؟",
             "هل يقلل تأخير الإقلاع والوصول من الرضا؟", "هل تؤثر مسافة الرحلة — وهل يعتمد ذلك على الدرجة؟",
             "أي مزيج من نوع السفر والدرجة هو الأسعد والأتعس؟"]
for i, q in enumerate(questions):
    x, y = 0.6 + (i % 2) * 6.15, 1.8 + (i // 2) * 1.65
    deck.card(s, x, y, 5.9, 1.35)
    deck.badge(s, x + 0.3, y + 0.38, f"Q{i + 1}", fill=SKY if i % 2 == 0 else NAVY, size=0.62, font_size=15)
    T(s, x + 1.15, y + 0.2, 4.55, 0.95, q, size=17, anchor=MSO_ANCHOR.MIDDLE)
deck.notes(s, "هذه الأسئلة الستة هي التي تنظّم التحليل في النوتبوك وفي لوحة البيانات.")

# ================================================================ 8. Divider
s = deck.dark_slide()
deck.shape(s, MSO_SHAPE.OVAL, 9.3, 1.4, 4.6, 4.6, NAVY_2)
T(s, 0.8, 2.5, 8.5, 0.5, "الجزء الثاني", size=17, color=LIGHT_MUTED, bold=True)
T(s, 0.8, 3.0, 8.5, 1.2, "النتائج المستخلصة من البيانات", size=46, color=WHITE, bold=True)
T(s, 0.8, 4.3, 8.5, 0.6, "كل نتيجة مدعومة باختبار إحصائي (Chi-square و Mann-Whitney والارتباط)", size=18,
  color=LIGHT_MUTED)
deck.notes(s, "ننتقل الآن إلى النتائج.")

# ================================================================ 9. Q1
s = deck.content_slide("معظم الركاب غير راضين", "السؤال الأول · الرضا العام")
counts = df["satisfaction"].value_counts()
cd = CategoryChartData()
cd.categories = ["راضٍ", "محايد أو غير راضٍ"]
cd.add_series("الركاب", [int(counts["Satisfied"]), int(counts["Neutral or Dissatisfied"])])
ch = deck.chart(s, XL_CHART_TYPE.DOUGHNUT, 0.6, 1.6, 6.0, 5.2, cd)
ch.has_title = False
ch.has_legend = True
ch.legend.position = XL_LEGEND_POSITION.BOTTOM
ch.legend.include_in_layout = False
ch.legend.font.size = Pt(14)
ch.font.name = FONT
for idx, colr in enumerate([SKY, ORANGE]):
    pt = ch.plots[0].series[0].points[idx]
    pt.format.fill.solid()
    pt.format.fill.fore_color.rgb = colr
    pt.format.line.color.rgb = WHITE
ch.plots[0].has_data_labels = True
dl = ch.plots[0].data_labels
dl.show_percentage, dl.show_value = True, False
dl.number_format, dl.number_format_is_linked = "0.0%", False
dl.font.size, dl.font.bold, dl.font.color.rgb = Pt(16), True, WHITE
T(s, 7.3, 2.0, 5.4, 1.2, "43.9%", size=72, color=SKY, bold=True, rtl=False, align=PP_ALIGN.RIGHT)
T(s, 7.3, 3.25, 5.4, 0.5, "من الركاب راضون", size=22, color=NAVY, bold=True)
T(s, 7.3, 4.0, 5.4, 2.0, [("14,573 راكبًا (56.1%) كانوا محايدين أو غير راضين.", {"space_after": 10}),
                          ("المجموعتان متوازنتان بشكل معقول، لذلك المقارنة بينهما موثوقة.", {})], size=17)
deck.notes(s, "إجابة السؤال الأول: 43.9% فقط من الركاب راضون. الأغلبية غير سعيدة.")

# ================================================================ 10. Q2 segments
s = deck.content_slide("الدرجة وغرض السفر يقسمان الركاب بوضوح", "السؤال الثاني · شرائح الركاب")
segs = [("درجة البيزنس", rate("class")["Business"]), ("إيكونومي بلس", rate("class")["Economy Plus"]),
        ("الإيكونومي", rate("class")["Economy"]), ("سفر عمل", rate("type_of_travel")["Business Travel"]),
        ("سفر شخصي", rate("type_of_travel")["Personal Travel"]), ("عميل دائم", rate("customer_type")["Loyal Customer"]),
        ("عميل غير دائم", rate("customer_type")["Disloyal Customer"]), ("أنثى", rate("gender")["Female"]),
        ("ذكر", rate("gender")["Male"])]
cd = CategoryChartData()
cd.categories = [n for n, _ in segs]
cd.add_series("نسبة الرضا", [round(v, 1) for _, v in segs])
ch = deck.chart(s, XL_CHART_TYPE.COLUMN_CLUSTERED, 0.6, 1.6, 8.4, 5.2, cd)
style_chart(ch, value_axis_max=100)
ch.category_axis.reverse_order = True
ch.plots[0].gap_width = 45
label_series(ch.plots[0])
color_series(ch, [SKY])
for i, (_, v) in enumerate(segs):
    p = ch.plots[0].series[0].points[i]
    p.format.fill.solid()
    p.format.fill.fore_color.rgb = SKY if v >= 43.9 else ORANGE
ch.category_axis.tick_labels.font.size = Pt(11)
deck.card(s, 9.4, 1.75, 3.35, 4.9)
T(s, 9.6, 1.95, 2.95, 4.6, [
    ("أقوى العوامل", {"bold": True, "color": NAVY, "size": 17, "space_after": 8}),
    ([("الدرجة ", {"bold": True}), ("— البيزنس 69.5% مقابل الإيكونومي 19.4%", {})], {"space_after": 8}),
    ([("غرض السفر ", {"bold": True}), ("— العمل 58.8% مقابل الشخصي 10.0%", {})], {"space_after": 8}),
    ([("الولاء ", {"bold": True}), ("— الدائم 48.1% مقابل غير الدائم 25.2%", {})], {"space_after": 8}),
    ([("النوع ", {"bold": True}), ("— لا يوجد فرق (p = 0.24)", {})], {"space_after": 8}),
    ("الأزرق = أعلى من المتوسط 43.9%", {"size": 12, "color": MUTED}),
], size=14)
deck.notes(s, "نسبة الرضا لكل شريحة. الدرجة وغرض السفر لهما أقوى ارتباط (Cramér's V بقيمة 0.50 و 0.45). "
              "النوع ليس له تأثير معنوي.")

# ================================================================ 11. Age
s = deck.content_slide("الركاب في منتصف العمر هم الأكثر رضا", "السؤال الثاني · العمر")
ar = rate("age_group", ["Child (<18)", "Young Adult (18-29)", "Adult (30-44)", "Middle-Aged (45-59)", "Senior (60+)"])
cd = CategoryChartData()
cd.categories = ["أطفال\nأقل من 18", "شباب\n18-29", "بالغون\n30-44", "منتصف العمر\n45-59", "كبار السن\n60+"]
cd.add_series("نسبة الرضا", [round(v, 1) for v in ar.values])
ch = deck.chart(s, XL_CHART_TYPE.LINE_MARKERS, 0.6, 1.6, 8.4, 5.2, cd)
style_chart(ch, value_axis_max=80)
ch.category_axis.reverse_order = True
ser = ch.plots[0].series[0]
ser.format.line.color.rgb = SKY
ser.format.line.width = Pt(3.5)
ser.smooth = False
ser.marker.size = 12
ser.marker.format.fill.solid()
ser.marker.format.fill.fore_color.rgb = SKY
ser.marker.format.line.color.rgb = WHITE
label_series(ch.plots[0], position=XL_LABEL_POSITION.ABOVE)
deck.card(s, 9.4, 1.75, 3.35, 4.9)
T(s, 9.6, 1.95, 2.95, 4.6, [
    ("57.7%", {"bold": True, "color": SKY, "size": 40, "rtl": False, "align": PP_ALIGN.RIGHT}),
    ("من فئة 45–59 سنة راضون", {"bold": True, "color": NAVY, "space_after": 14}),
    ("الأطفال (18%) والشباب (36%) وكبار السن (27%) هم الأقل رضا — وهي فئات تسافر غالبًا لأسباب شخصية "
     "على الدرجة الاقتصادية.", {"space_after": 10}),
    ("وسيط العمر: 43 للراضين مقابل 37 لغير الراضين (Mann-Whitney p < 0.001)", {"size": 12, "color": MUTED}),
], size=14)
deck.notes(s, "يرتفع الرضا مع العمر حتى حوالي الستين، ثم ينخفض مرة أخرى.")

# ================================================================ 12. Q3 services
s = deck.content_slide("الصعود أونلاين والترفيه والواي فاي تصنع الفرق", "السؤال الثالث · الخدمات")
means = df.groupby("satisfaction")[SERVICES].mean().T
means["gap"] = means["Satisfied"] - means["Neutral or Dissatisfied"]
means = means.sort_values("gap")
arabic_services = {"inflight_wifi_service": "الواي فاي", "departure_arrival_time_convenient": "ملاءمة المواعيد",
                   "ease_of_online_booking": "الحجز أونلاين", "gate_location": "موقع البوابة",
                   "food_and_drink": "الطعام والشراب", "online_boarding": "الصعود أونلاين",
                   "seat_comfort": "راحة المقعد", "inflight_entertainment": "الترفيه",
                   "on_board_service": "الخدمة على المتن", "leg_room_service": "مساحة الأرجل",
                   "baggage_handling": "الأمتعة", "checkin_service": "تسجيل الوصول",
                   "inflight_service": "الخدمة أثناء الرحلة", "cleanliness": "النظافة"}
cd = CategoryChartData()
cd.categories = [arabic_services[c] for c in means.index]
cd.add_series("غير راضٍ", [round(v, 2) for v in means["Neutral or Dissatisfied"]])
cd.add_series("راضٍ", [round(v, 2) for v in means["Satisfied"]])
ch = deck.chart(s, XL_CHART_TYPE.BAR_CLUSTERED, 0.6, 1.5, 8.0, 5.4, cd)
style_chart(ch, legend=True, value_axis_max=5, number_format="0", font_size=11)
ch.value_axis.major_unit = 1
ch.value_axis.reverse_order = True          # bars grow right-to-left
ch.category_axis.tick_labels.font.size = Pt(11)
ch.plots[0].gap_width = 40
ch.plots[0].overlap = 0
color_series(ch, [ORANGE, SKY])
deck.card(s, 9.0, 1.75, 3.75, 4.9)
T(s, 9.2, 1.95, 3.35, 4.6, [
    ("أكبر فجوات التقييم", {"bold": True, "color": NAVY, "size": 17, "space_after": 8}),
    ([("+1.44 ", {"bold": True, "color": SKY}), ("الصعود أونلاين", {})], {"space_after": 4}),
    ([("+1.08 ", {"bold": True, "color": SKY}), ("الترفيه", {})], {"space_after": 4}),
    ([("+1.00 ", {"bold": True, "color": SKY}), ("الواي فاي", {})], {"space_after": 4}),
    ([("+0.92 ", {"bold": True, "color": SKY}), ("راحة المقعد", {})], {"space_after": 12}),
    ([("الواي فاي هو الأقل تقييمًا (2.81) ", {"bold": True, "color": ORANGE}),
      ("— ومع ذلك 99% ممن قيّموه بـ 5 راضون.", {})], {"space_after": 10}),
    ("موقع البوابة وملاءمة المواعيد لا علاقة لهما بالرضا.", {"size": 13, "color": MUTED}),
], size=14)
deck.notes(s, "متوسط تقييم كل خدمة للراضين مقابل غير الراضين. أكبر الفجوات توضح الخدمات التي تفصل الراكب السعيد عن "
              "غير السعيد. الصعود أونلاين له أقوى ارتباط بقيمة 0.59.")

# ================================================================ 13. Q4 delays
s = deck.content_slide("التأخير يقلل الرضا — لكن بدرجة محدودة", "السؤال الرابع · التأخير")
dr = rate("delay_group", ["On time", "1-15 min", "16-60 min", "> 60 min"])
cd = CategoryChartData()
cd.categories = ["في الموعد", "1–15 دقيقة", "16–60 دقيقة", "أكثر من 60 دقيقة"]
cd.add_series("نسبة الرضا", [round(v, 1) for v in dr.values])
ch = deck.chart(s, XL_CHART_TYPE.COLUMN_CLUSTERED, 0.6, 1.6, 7.6, 5.2, cd)
style_chart(ch, value_axis_max=60)
ch.category_axis.reverse_order = True
ch.plots[0].gap_width = 60
label_series(ch.plots[0])
color_series(ch, [SKY])
for i, (num, lab) in enumerate([("12.7 نقطة", "انخفاض من الوصول في الموعد إلى تأخير أكثر من ساعة"),
                                ("56.3%", "من الرحلات تصل في موعدها"),
                                ("V = 0.10", "ارتباط ضعيف جدًا (p < 0.001)")]):
    y = 1.75 + i * 1.65
    deck.card(s, 8.7, y, 4.05, 1.4)
    T(s, 8.95, y + 0.15, 3.6, 0.6, num, size=28, color=ORANGE if i == 0 else NAVY, bold=True)
    T(s, 8.95, y + 0.82, 3.6, 0.5, lab, size=13, color=MUTED)
deck.notes(s, "ينخفض الرضا تدريجيًا من 47.9% عند الوصول في الموعد إلى 35.2% عند التأخير أكثر من ساعة. الفرق معنوي، "
              "لكنه أضعف بكثير من تأثير الدرجة أو جودة الخدمة.")

# ================================================================ 14. Q5 distance x class
s = deck.content_slide("المسافة تؤثر فقط من خلال الدرجة", "السؤال الخامس · مسافة الرحلة")
pv = (df.pivot_table(index="distance_group", columns="class", values="is_satisfied", aggfunc="mean") * 100).reindex(
    ["Short-haul (<800 mi)", "Medium-haul (800-2000 mi)", "Long-haul (>2000 mi)"])
cd = CategoryChartData()
cd.categories = ["قصيرة\nأقل من 800 ميل", "متوسطة\n800–2,000 ميل", "طويلة\nأكثر من 2,000 ميل"]
for cls, name in [("Economy", "الإيكونومي"), ("Economy Plus", "إيكونومي بلس"), ("Business", "البيزنس")]:
    cd.add_series(name, [round(v, 1) for v in pv[cls].values])
ch = deck.chart(s, XL_CHART_TYPE.COLUMN_CLUSTERED, 0.6, 1.5, 8.0, 5.3, cd)
style_chart(ch, legend=True, value_axis_max=100)
ch.category_axis.reverse_order = True
ch.plots[0].gap_width = 60
ch.plots[0].overlap = -5
label_series(ch.plots[0], number_format='0"%"', size=11)
color_series(ch, [ORANGE, YELLOW, SKY])
deck.card(s, 9.0, 1.75, 3.75, 4.9)
T(s, 9.2, 1.95, 3.35, 4.6, [
    ("عامل خفي", {"bold": True, "color": NAVY, "size": 17, "space_after": 8}),
    ("بشكل عام يبدو ركاب الرحلات الطويلة أكثر رضا (70% مقابل 34% للقصيرة).", {"space_after": 10}),
    ("لكن الرحلات الطويلة غالبًا على درجة البيزنس (الوسيط 1,590 ميل مقابل حوالي 600 ميل في الإيكونومي).",
     {"space_after": 10}),
    ([("داخل الإيكونومي، المسافة لا تغيّر شيئًا تقريبًا (17–21%). ", {"bold": True}),
      ("فقط البيزنس يتحسن في الرحلات الطويلة (من 60% إلى 77%).", {})], {}),
], size=14)
deck.notes(s, "هذا مثال جيد على أهمية التحليل متعدد المتغيرات: تأثير المسافة يختفي عندما نثبّت الدرجة.")

# ================================================================ 15. Q6 combination
s = deck.content_slide("المسافرون لأسباب شخصية غير سعداء في كل الدرجات", "السؤال السادس · نوع السفر × الدرجة")
tc = df.pivot_table(index="type_of_travel", columns="class", values="is_satisfied", aggfunc="mean") * 100
tn = df.pivot_table(index="type_of_travel", columns="class", values="id", aggfunc="count")
classes = [("Economy", "الإيكونومي"), ("Economy Plus", "إيكونومي بلس"), ("Business", "البيزنس")]
travel = [("Business Travel", "سفر عمل"), ("Personal Travel", "سفر شخصي")]
x0, y0, cw, chh = 3.1, 2.15, 2.15, 1.75
for j, (_, cname) in enumerate(classes):
    T(s, x0 + j * (cw + 0.12), y0 - 0.5, cw, 0.4, cname, size=16, color=NAVY, bold=True, align=PP_ALIGN.CENTER)
for i, (tr, tname) in enumerate(travel):
    T(s, 0.6, y0 + i * (chh + 0.12) + 0.55, 2.4, 0.6, tname, size=17, color=NAVY, bold=True)
    for j, (cls, _) in enumerate(classes):
        v = tc.loc[tr, cls]
        shade = SKY if v > 60 else (RGBColor(0x9C, 0xC3, 0xEE) if v > 25 else ICE)
        tx = WHITE if v > 60 else NAVY
        x, y = x0 + j * (cw + 0.12), y0 + i * (chh + 0.12)
        deck.card(s, x, y, cw, chh, fill=shade)
        T(s, x, y + 0.35, cw, 0.7, f"{v:.1f}%", size=32, color=tx, bold=True, rtl=False, align=PP_ALIGN.CENTER)
        T(s, x, y + 1.1, cw, 0.4, f"n = {tn.loc[tr, cls]:,}", size=12, color=tx, rtl=False, align=PP_ALIGN.CENTER)
T(s, 0.6, 6.0, 9.2, 0.6, "نسبة الركاب الراضين لكل مزيج (اللون الأغمق = رضا أعلى)", size=12, color=MUTED)
deck.card(s, 10.0, 1.65, 2.75, 4.7)
T(s, 10.15, 1.85, 2.45, 4.4, [
    ("الأسعد", {"bold": True, "color": SKY, "size": 17}),
    ("المسافرون للعمل على درجة البيزنس — 72%", {"space_after": 14}),
    ("الأتعس", {"bold": True, "color": ORANGE, "size": 17}),
    ("المسافرون لأسباب شخصية — حوالي 9–10% في كل الدرجات، حتى البيزنس", {"space_after": 14}),
    ("الولاء لا يساعدهم: المسافرون الدائمون لأسباب شخصية راضون بنسبة 10% فقط.", {"size": 13, "color": MUTED}),
], size=14)
deck.notes(s, "شريحة العمل الأساسية تحصل على خدمة جيدة. المسافرون لأسباب شخصية هم المشكلة الكبرى بغض النظر عن الدرجة.")

# ================================================================ 16. Drivers
s = deck.content_slide("ما الذي يؤثر على الرضا أكثر؟", "ملخص العوامل")
drivers = [("الدرجة", cramers_v("class")), ("غرض السفر", cramers_v("type_of_travel")),
           ("الفئة العمرية", cramers_v("age_group")), ("نوع العميل", cramers_v("customer_type")),
           ("تأخير الوصول", cramers_v("delay_group")), ("النوع", cramers_v("gender"))][::-1]
cd = CategoryChartData()
cd.categories = [d for d, _ in drivers]
cd.add_series("Cramér's V", [round(v, 2) for _, v in drivers])
ch = deck.chart(s, XL_CHART_TYPE.BAR_CLUSTERED, 0.6, 1.6, 7.6, 5.0, cd)
style_chart(ch, value_axis_max=0.6, number_format="0.0")
ch.value_axis.reverse_order = True
ch.plots[0].gap_width = 50
label_series(ch.plots[0], number_format="0.00")
color_series(ch, [SKY])
ch.category_axis.tick_labels.font.size = Pt(13)
deck.card(s, 8.7, 1.75, 4.05, 4.9)
T(s, 8.9, 1.95, 3.65, 4.6, [
    ("كيف تقرأ الرسم", {"bold": True, "color": NAVY, "size": 17, "space_after": 8}),
    ("مقياس Cramér's V يوضح قوة ارتباط كل عامل بالرضا (0 = لا يوجد، 0.5 = قوي).", {"space_after": 12}),
    ([("جودة الخدمة ", {"bold": True}), ("(الصعود أونلاين r = 0.59) قوية بقدر الدرجة.", {})], {"space_after": 12}),
    ([("التأخير والنوع ", {"bold": True}), ("تأثيرهما ضئيل.", {})], {}),
], size=14)
deck.notes(s, "ترتيب كل العوامل حسب قوة ارتباطها بالرضا.")

# ================================================================ 17. Recommendations
s = deck.content_slide("توصيات لشركة الطيران", "الخطوات القادمة")
recs = [("إصلاح الرحلة الرقمية", "الواي فاي والصعود والحجز أونلاين: الأقل تقييمًا ومن أقوى العوامل"),
        ("تطوير الدرجة الاقتصادية", "أقل من راكب من كل 5 في الإيكونومي راضٍ — المقاعد والترفيه ومساحة الأرجل"),
        ("كسب المسافرين لأسباب شخصية", "حوالي 10% فقط راضون: خدمات عائلية، أسعار مرنة، ترفيه"),
        ("استعادة العملاء غير الدائمين", "25% راضون — عروض ولاء موجّهة"),
        ("تقليل التأخير", "تأثير أصغر، لكن الرضا ينخفض مع كل فئة تأخير")]
for i, (head, desc) in enumerate(recs):
    y = 1.7 + i * 1.03
    deck.card(s, 0.6, y, 12.15, 0.88)
    deck.badge(s, 0.85, y + 0.16, str(i + 1), fill=SKY if i < 3 else NAVY)
    T(s, 1.65, y + 0.12, 3.9, 0.65, head, size=18, color=NAVY, bold=True, anchor=MSO_ANCHOR.MIDDLE)
    T(s, 5.6, y + 0.12, 6.9, 0.65, desc, size=15, anchor=MSO_ANCHOR.MIDDLE)
deck.notes(s, "خمس توصيات مرتبة حسب التأثير المتوقع.")

# ================================================================ 18. Conclusion
s = deck.dark_slide()
T(s, 0.8, 0.7, 11.8, 0.5, "الخلاصة", size=17, color=LIGHT_MUTED, bold=True)
T(s, 0.8, 1.2, 11.8, 1.0, "الرضا يُكتسب داخل الطائرة وعلى الإنترنت — لا بالمسافة ولا بالنوع", size=32,
  color=WHITE, bold=True)
for i, (num, lab) in enumerate([("43.9%", "راضون بشكل عام — هناك مجال كبير للتحسن"),
                                ("72% مقابل 10%", "المسافرون للعمل على البيزنس مقابل المسافرين لأسباب شخصية"),
                                ("+1.44", "الصعود أونلاين: أكبر فجوة في الخدمات")]):
    x = 0.8 + i * 4.0
    deck.shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, x, 3.0, 3.7, 2.4, NAVY_2)
    T(s, x + 0.3, 3.3, 3.1, 0.9, num, size=34, color=WHITE, bold=True)
    T(s, x + 0.3, 4.25, 3.1, 1.0, lab, size=15, color=LIGHT_MUTED)
T(s, 0.8, 5.9, 11.8, 0.8, "تم تنظيف البيانات (تقييمات «لا ينطبق»، التأخير الناقص، التسميات غير المتسقة)، وتحليل أكثر من "
                         "10 متغيرات بستة أنواع من الرسوم، والتحقق من كل نتيجة باختبار إحصائي.", size=15,
  color=LIGHT_MUTED)
deck.notes(s, "الختام: هذه هي الأرقام الثلاثة التي يجب تذكرها.")

# ================================================================ 19. Thank you
s = deck.dark_slide()
deck.shape(s, MSO_SHAPE.OVAL, 8.9, 1.2, 5.0, 5.0, NAVY_2)
deck.shape(s, MSO_SHAPE.OVAL, 10.6, 4.4, 2.2, 2.2, SKY)
T(s, 0.8, 2.3, 8, 1.2, "شكرًا لكم", size=60, color=WHITE, bold=True)
T(s, 0.8, 3.6, 8, 0.6, "الأسئلة والنقاش", size=24, color=LIGHT_MUTED)
T(s, 0.8, 4.8, 8, 1.2, [("github.com/hahmed1936/Mid-Project", {"rtl": False, "align": PP_ALIGN.RIGHT,
                                                               "space_after": 6}),
                        ("لوحة البيانات التفاعلية: Streamlit (الرابط في ملف README)", {})], size=16, color=WHITE)

print("saved", OUT, deck.save(OUT), "slides")
