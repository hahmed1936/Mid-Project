"""Shared helpers for building right-to-left (Arabic) slide decks with python-pptx."""
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_LEGEND_POSITION, XL_LABEL_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

NAVY = RGBColor(0x0B, 0x1F, 0x3A)
NAVY_2 = RGBColor(0x15, 0x30, 0x55)
SKY = RGBColor(0x2A, 0x78, 0xD6)
ORANGE = RGBColor(0xEB, 0x68, 0x34)
YELLOW = RGBColor(0xED, 0xA1, 0x00)
ICE = RGBColor(0xE8, 0xF1, 0xFB)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
INK = RGBColor(0x1B, 0x1F, 0x24)
MUTED = RGBColor(0x5B, 0x63, 0x6E)
LIGHT_MUTED = RGBColor(0xB8, 0xC7, 0xDB)
GRID = RGBColor(0xDD, 0xE3, 0xEA)
CODE_BG = RGBColor(0x14, 0x1C, 0x2B)
CODE_TEXT = RGBColor(0xE6, 0xED, 0xF7)
CODE_COMMENT = RGBColor(0x8F, 0xD6, 0x9B)
FONT = "Arial"            # renders Arabic correctly in every Office install
CODE_FONT = "Courier New"
SLIDE_W, SLIDE_H = 13.333, 7.5


class ArabicDeck:
    """A 16:9 deck whose helpers mirror positions horizontally (right-to-left layout) and write RTL text."""

    def __init__(self, footer, mirror=True):
        self.prs = Presentation()
        self.prs.slide_width, self.prs.slide_height = Inches(SLIDE_W), Inches(SLIDE_H)
        self.blank = self.prs.slide_layouts[6]
        self.footer = footer
        self.mirror = mirror
        self.slide_count = 0

    # ------------------------------------------------------------ geometry
    def mx(self, x, w):
        """Mirror an x position so a layout designed left-to-right reads right-to-left."""
        return SLIDE_W - x - w if self.mirror else x

    # ------------------------------------------------------------ primitives
    def text(self, slide, x, y, w, h, content, size=16, color=INK, bold=False, rtl=True, font=FONT, align=None,
             anchor=MSO_ANCHOR.TOP, italic=False):
        """Text box. `content` is a string or a list of (text, options) paragraphs; runs may be (text, options)."""
        box = slide.shapes.add_textbox(Inches(self.mx(x, w)), Inches(y), Inches(w), Inches(h))
        tf = box.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        tf.vertical_anchor = anchor
        paragraphs = content if isinstance(content, list) else [(content, {})]
        for i, (txt, opts) in enumerate(paragraphs):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            para_rtl = opts.get("rtl", rtl)
            p.alignment = opts.get("align", align or (PP_ALIGN.RIGHT if para_rtl else PP_ALIGN.LEFT))
            if para_rtl:
                p._p.get_or_add_pPr().set("rtl", "1")
            p.space_after = Pt(opts.get("space_after", 0))
            runs = txt if isinstance(txt, list) else [(txt, {})]
            for rtxt, ropts in runs:
                r = p.add_run()
                r.text = rtxt
                f = r.font
                f.name = ropts.get("font", opts.get("font", font))
                f.size = Pt(ropts.get("size", opts.get("size", size)))
                f.bold = ropts.get("bold", opts.get("bold", bold))
                f.italic = ropts.get("italic", opts.get("italic", italic))
                f.color.rgb = ropts.get("color", opts.get("color", color))
                rpr = r._r.get_or_add_rPr()          # complex-script (Arabic) font slot
                cs = rpr.find(qn("a:cs"))
                if cs is None:
                    cs = rpr.makeelement(qn("a:cs"), {})
                    rpr.append(cs)
                cs.set("typeface", f.name)
        return box

    def shape(self, slide, kind, x, y, w, h, fill):
        s = slide.shapes.add_shape(kind, Inches(self.mx(x, w)), Inches(y), Inches(w), Inches(h))
        s.fill.solid()
        s.fill.fore_color.rgb = fill
        s.line.fill.background()
        s.shadow.inherit = False
        if kind == MSO_SHAPE.ROUNDED_RECTANGLE:
            s.adjustments[0] = 0.08
        return s

    def card(self, slide, x, y, w, h, fill=ICE):
        return self.shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h, fill)

    def badge(self, slide, x, y, label, fill=SKY, size=0.55, font_size=16):
        """The decks' motif: a numbered circle."""
        c = self.shape(slide, MSO_SHAPE.OVAL, x, y, size, size, fill)
        tf = c.text_frame
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        r.text = label
        r.font.size, r.font.bold, r.font.color.rgb, r.font.name = Pt(font_size), True, WHITE, FONT
        return c

    def chart(self, slide, chart_type, x, y, w, h, data):
        return slide.shapes.add_chart(chart_type, Inches(self.mx(x, w)), Inches(y), Inches(w), Inches(h), data).chart

    def table(self, slide, rows, cols, x, y, w, h):
        return slide.shapes.add_table(rows, cols, Inches(self.mx(x, w)), Inches(y), Inches(w), Inches(h)).table

    def code_block(self, slide, x, y, w, h, code, max_size=12, min_size=8.5):
        """Dark code panel; font shrinks so the longest line and all lines fit. Comments are green."""
        self.card(slide, x, y, w, h, fill=CODE_BG)
        lines = code.split("\n")
        longest = max(len(line) for line in lines)
        by_width = (w - 0.5) * 72 / (0.6 * max(longest, 1))
        by_height = (h - 0.35) * 72 / (1.2 * len(lines))
        size = max(min_size, min(max_size, by_width, by_height))
        paragraphs = [(line if line else " ",
                       {"color": CODE_COMMENT if line.strip().startswith("#") else CODE_TEXT})
                      for line in lines]
        box = self.text(slide, x + 0.25, y + 0.18, w - 0.5, h - 0.3, paragraphs, size=round(size, 1), rtl=False,
                        font=CODE_FONT)
        return box

    @staticmethod
    def notes(slide, script):
        slide.notes_slide.notes_text_frame.text = script

    # ------------------------------------------------------------ slide frames
    def content_slide(self, title, kicker=None):
        s = self.prs.slides.add_slide(self.blank)
        s.background.fill.solid()
        s.background.fill.fore_color.rgb = WHITE
        self.slide_count += 1
        if kicker:
            self.text(s, 0.6, 0.38, 12.1, 0.3, kicker, size=13, color=SKY, bold=True)
        self.text(s, 0.6, 0.68, 12.1, 0.8, title, size=30, color=NAVY, bold=True)
        self.text(s, 0.6, 7.0, 8, 0.3, self.footer, size=10, color=MUTED)
        self.text(s, 12.2, 7.0, 0.55, 0.3, str(self.slide_count), size=10, color=MUTED, rtl=False,
                  align=PP_ALIGN.LEFT if self.mirror else PP_ALIGN.RIGHT)
        return s

    def dark_slide(self):
        s = self.prs.slides.add_slide(self.blank)
        s.background.fill.solid()
        s.background.fill.fore_color.rgb = NAVY
        self.slide_count += 1
        return s

    def save(self, path):
        self.prs.save(path)
        return len(self.prs.slides)


# ---------------------------------------------------------------- native chart styling
def style_chart(chart, legend=False, value_axis_max=None, number_format='0"%"', font_size=12, gridlines=True):
    chart.has_title = False
    chart.font.name = FONT
    chart.font.size = Pt(font_size)
    chart.font.color.rgb = MUTED
    chart.has_legend = legend
    if legend:
        chart.legend.position = XL_LEGEND_POSITION.TOP
        chart.legend.include_in_layout = False
        chart.legend.font.size = Pt(12)
        chart.legend.font.color.rgb = INK
    va = chart.value_axis
    va.has_major_gridlines = gridlines
    if gridlines:
        va.major_gridlines.format.line.color.rgb = GRID
    va.format.line.fill.background()
    va.tick_labels.font.size = Pt(11)
    va.tick_labels.number_format = number_format
    va.tick_labels.number_format_is_linked = False
    if value_axis_max is not None:
        va.maximum_scale = value_axis_max
        va.minimum_scale = 0
    ca = chart.category_axis
    ca.format.line.color.rgb = GRID
    ca.tick_labels.font.size = Pt(12)
    ca.tick_labels.font.color.rgb = INK
    ca.has_major_gridlines = False


def label_series(plot, number_format='0.0"%"', size=12, color=INK, position=XL_LABEL_POSITION.OUTSIDE_END):
    plot.has_data_labels = True
    dl = plot.data_labels
    dl.number_format = number_format
    dl.number_format_is_linked = False
    dl.font.size = Pt(size)
    dl.font.bold = True
    dl.font.color.rgb = color
    if position is not None:
        dl.position = position


def color_series(chart, colors):
    for series, c in zip(chart.plots[0].series, colors):
        series.format.fill.solid()
        series.format.fill.fore_color.rgb = c
        series.format.line.color.rgb = WHITE
        series.format.line.width = Pt(1.5)
