"""Build the Arabic deck used while recording the code-walkthrough video (python-pptx, right-to-left text)."""
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

PROJECT = Path(__file__).parent
OUT = PROJECT / "Airline_Satisfaction_Video_AR.pptx"

NAVY = RGBColor(0x0B, 0x1F, 0x3A)
NAVY_2 = RGBColor(0x15, 0x30, 0x55)
SKY = RGBColor(0x2A, 0x78, 0xD6)
ORANGE = RGBColor(0xEB, 0x68, 0x34)
ICE = RGBColor(0xE8, 0xF1, 0xFB)
CODE_BG = RGBColor(0x14, 0x1C, 0x2B)
CODE_TEXT = RGBColor(0xE6, 0xED, 0xF7)
CODE_GREEN = RGBColor(0x8F, 0xD6, 0x9B)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
INK = RGBColor(0x1B, 0x1F, 0x24)
MUTED = RGBColor(0x5B, 0x63, 0x6E)
LIGHT_MUTED = RGBColor(0xB8, 0xC7, 0xDB)
FONT = "Arial"          # renders Arabic correctly in every Office install
CODE_FONT = "Courier New"

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
BLANK = prs.slide_layouts[6]
slide_counter = [0]


def set_rtl(paragraph):
    """Mark a paragraph as right-to-left so Arabic punctuation and mixed English words order correctly."""
    paragraph._p.get_or_add_pPr().set("rtl", "1")


def text(slide, x, y, w, h, content, size=18, color=INK, bold=False, rtl=True, font=FONT,
         align=None, anchor=MSO_ANCHOR.TOP):
    """Text box; `content` is a string or a list of (text, options) paragraphs. Arabic is right-aligned by default."""
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    paragraphs = content if isinstance(content, list) else [(content, {})]
    for i, (txt, opts) in enumerate(paragraphs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para_rtl = opts.get("rtl", rtl)
        p.alignment = opts.get("align", align or (PP_ALIGN.RIGHT if para_rtl else PP_ALIGN.LEFT))
        if para_rtl:
            set_rtl(p)
        p.space_after = Pt(opts.get("space_after", 0))
        runs = txt if isinstance(txt, list) else [(txt, {})]
        for rtxt, ropts in runs:
            r = p.add_run()
            r.text = rtxt
            f = r.font
            f.name = ropts.get("font", opts.get("font", font))
            f.size = Pt(ropts.get("size", opts.get("size", size)))
            f.bold = ropts.get("bold", opts.get("bold", bold))
            f.color.rgb = ropts.get("color", opts.get("color", color))
            # make sure Arabic glyphs use the same font (complex-script font slot)
            rpr = r._r.get_or_add_rPr()
            cs = rpr.find(qn("a:cs"))
            if cs is None:
                cs = rpr.makeelement(qn("a:cs"), {})
                rpr.append(cs)
            cs.set("typeface", f.name)
    return box


def shape(slide, kind, x, y, w, h, fill):
    s = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    s.line.fill.background()
    s.shadow.inherit = False
    if kind == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = 0.06
    return s


def badge(slide, x, y, label, fill=SKY, size=0.55, font_size=16):
    c = shape(slide, MSO_SHAPE.OVAL, x, y, size, size, fill)
    tf = c.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = label
    r.font.size, r.font.bold, r.font.color.rgb, r.font.name = Pt(font_size), True, WHITE, FONT


def code_block(slide, x, y, w, h, code, size=13):
    """Dark code panel with left-to-right monospace code; comment lines are green."""
    shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h, CODE_BG)
    lines = [(line if line else " ", {"color": CODE_GREEN if line.strip().startswith("#") else CODE_TEXT})
             for line in code.strip("\n").split("\n")]
    text(slide, x + 0.3, y + 0.25, w - 0.6, h - 0.4, lines, size=size, rtl=False, font=CODE_FONT)


def notes(slide, script):
    slide.notes_slide.notes_text_frame.text = script


def content_slide(title, kicker):
    s = prs.slides.add_slide(BLANK)
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = WHITE
    slide_counter[0] += 1
    text(s, 0.6, 0.38, 12.13, 0.35, kicker, size=14, color=SKY, bold=True)
    text(s, 0.6, 0.75, 12.13, 0.8, title, size=32, color=NAVY, bold=True)
    text(s, 4.73, 7.0, 8.0, 0.3, "شرح كود مشروع رضا ركاب الطيران", size=10, color=MUTED)
    text(s, 0.6, 7.0, 0.6, 0.3, str(slide_counter[0] + 1), size=10, color=MUTED, rtl=False)
    return s


def dark_slide():
    s = prs.slides.add_slide(BLANK)
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = NAVY
    slide_counter[0] += 1
    return s


def card(slide, x, y, w, h, fill=ICE):
    return shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h, fill)


# ================================================================ 1. Title
s = prs.slides.add_slide(BLANK)
s.background.fill.solid()
s.background.fill.fore_color.rgb = NAVY
shape(s, MSO_SHAPE.OVAL, -1.6, -1.2, 6.4, 6.4, NAVY_2)
shape(s, MSO_SHAPE.OVAL, 0.6, 3.9, 3.4, 3.4, SKY)
text(s, 0.6, 5.0, 3.4, 0.7, "Python", size=34, color=WHITE, bold=True, rtl=False, align=PP_ALIGN.CENTER)
text(s, 0.6, 5.75, 3.4, 0.4, "Jupyter Notebook", size=14, color=WHITE, rtl=False, align=PP_ALIGN.CENTER)
text(s, 4.5, 1.3, 8.2, 0.4, "المشروع المرحلي  ·  تحليل البيانات", size=16, color=LIGHT_MUTED, bold=True)
text(s, 4.5, 1.9, 8.2, 2.0, "شرح كود مشروع\nرضا ركاب الطيران", size=50, color=WHITE, bold=True)
text(s, 4.5, 4.1, 8.2, 0.8, "من البيانات الخام إلى النتائج — خطوة بخطوة", size=22, color=LIGHT_MUTED)
text(s, 4.5, 5.6, 8.2, 0.4, "هشام محمد", size=20, color=WHITE, bold=True)
text(s, 4.5, 6.1, 8.2, 0.4, "Epsilon AI  ·  برنامج علوم البيانات", size=14, color=LIGHT_MUTED)
notes(s, "السلام عليكم، في الفيديو ده هشرح الكود الخاص بمشروع تحليل رضا ركاب الطيران، من أول تحميل البيانات "
         "وتقييم جودتها، مرورًا بالتنظيف، وبعدين التحليل الاستكشافي والاختبارات الإحصائية، لحد النتائج النهائية.")

# ================================================================ 2. Agenda
s = content_slide("خطة الفيديو", "المحتوى")
steps = [("البيانات والأسئلة", "المجال، الأعمدة، والأسئلة البحثية الستة"),
         ("المكتبات والإعدادات", "الاستيراد والألوان الثابتة للرسوم"),
         ("تقييم جودة البيانات", "اكتشاف المشاكل قبل التنظيف"),
         ("تنظيف البيانات", "Define → Code → Test لكل مشكلة"),
         ("التحليل الاستكشافي", "أحادي المتغير، ثنائي ومتعدد المتغيرات، واختبارات إحصائية"),
         ("النتائج والخلاصة", "الإجابة على الأسئلة والتوصيات")]
for i, (head, sub) in enumerate(steps):
    y = 1.7 + i * 0.86
    badge(s, 12.15, y, str(i + 1), fill=SKY if i % 2 == 0 else NAVY)
    text(s, 4.0, y - 0.02, 7.9, 0.4, head, size=21, color=NAVY, bold=True)
    text(s, 4.0, y + 0.38, 7.9, 0.35, sub, size=14, color=MUTED)
shape(s, MSO_SHAPE.OVAL, 0.6, 2.2, 2.8, 2.8, ICE)
text(s, 0.6, 2.95, 2.8, 0.7, "89", size=48, color=SKY, bold=True, rtl=False, align=PP_ALIGN.CENTER)
text(s, 0.6, 3.8, 2.8, 0.4, "خلية في النوتبوك", size=15, color=NAVY, align=PP_ALIGN.CENTER)
notes(s, "الفيديو مقسم لست أجزاء بنفس ترتيب النوتبوك. النوتبوك فيه 89 خلية وبيشتغل من أوله لآخره من غير أي أخطاء.")

# ================================================================ 3. Data & questions
s = content_slide("البيانات والأسئلة البحثية", "الجزء الأول — البيانات")
facts = [("25,976", "راكب"), ("25", "عمود"), ("14", "خدمة متقيّمة من 1 لـ 5"), ("2", "فئة للرضا")]
for i, (num, lab) in enumerate(facts):
    x = 9.95 - (i % 2) * 2.95
    y = 1.75 + (i // 2) * 2.45
    card(s, x, y, 2.75, 2.15)
    text(s, x, y + 0.35, 2.75, 0.9, num, size=40, color=NAVY, bold=True, rtl=False, align=PP_ALIGN.CENTER)
    text(s, x + 0.15, y + 1.35, 2.45, 0.6, lab, size=15, color=MUTED, align=PP_ALIGN.CENTER)
questions = ["ما نسبة الركاب الراضين؟", "ما الشرائح الأكثر والأقل رضا؟", "ما الخدمات الأكثر تأثيرًا على الرضا؟",
             "هل التأخير يقلل الرضا؟", "هل المسافة تؤثر؟ وهل يعتمد ذلك على الدرجة؟",
             "أي مزيج من نوع السفر والدرجة هو الأسعد / الأتعس؟"]
text(s, 0.6, 1.75, 6.1, 0.5, "الأسئلة البحثية", size=20, color=NAVY, bold=True)
text(s, 0.6, 2.35, 6.1, 4.3, [([(f"Q{i + 1}   ", {"bold": True, "color": SKY}), (q, {})], {"space_after": 10})
                              for i, q in enumerate(questions)], size=16)
notes(s, "البيانات عبارة عن استبيان رضا ركاب طيران: حوالي 26 ألف راكب و25 عمود. كل راكب عنده بيانات شخصية زي "
         "النوع والسن والولاء، وبيانات الرحلة زي نوع السفر والدرجة والمسافة والتأخير، وتقييم لـ 14 خدمة، "
         "وفي الآخر عمود الرضا وهو الهدف. وحددت ست أسئلة هجاوب عليها في باقي التحليل.")

# ================================================================ 4. Imports
s = content_slide("المكتبات والإعدادات", "الجزء الثاني — الإعداد")
code_block(s, 0.6, 1.7, 7.4, 4.9, """
# Core data handling
import numpy as np
import pandas as pd

# Visualisation
import matplotlib.pyplot as plt
import seaborn as sns

# Statistics
from scipy import stats

# Fixed colours for every chart
SATISFIED_COLOR = "#2a78d6"     # blue
DISSATISFIED_COLOR = "#eb6834"  # orange
DATA_PATH = "Data.csv"
""")
text(s, 8.4, 1.75, 4.33, 4.9, [
    ("ليه المكتبات دي؟", {"bold": True, "color": NAVY, "size": 20, "space_after": 10}),
    ([("pandas / numpy", {"bold": True, "color": SKY}), (" للتعامل مع البيانات", {})], {"space_after": 8}),
    ([("matplotlib / seaborn", {"bold": True, "color": SKY}), (" للرسوم البيانية", {})], {"space_after": 8}),
    ([("scipy.stats", {"bold": True, "color": SKY}), (" للاختبارات الإحصائية", {})], {"space_after": 16}),
    ("ثبّتت لون لكل فئة: الأزرق = راضي، البرتقالي = غير راضي — نفس اللون في كل الرسوم علشان القارئ ما يتلخبطش.",
     {"color": MUTED, "size": 15}),
], size=16)
notes(s, "في البداية بستورد المكتبات: pandas وnumpy للبيانات، matplotlib وseaborn للرسم، وscipy للاختبارات "
         "الإحصائية. وكمان بثبّت ألوان: الأزرق للراضي والبرتقالي لغير الراضي، وده بيخلي كل الرسوم متسقة.")

# ================================================================ 5. Assessment
s = content_slide("تقييم جودة البيانات", "الجزء الثالث — التقييم")
code_block(s, 0.6, 1.7, 7.4, 4.9, """
raw_df = pd.read_csv(DATA_PATH)

def quality_report(df):
    \"\"\"Per-column quality summary.\"\"\"
    return pd.DataFrame({
        "dtype": df.dtypes.astype(str),
        "missing": df.isna().sum(),
        "unique": df.nunique(),
    })

raw_df.duplicated().sum()           # duplicates
(raw_df[service_columns] == 0).sum() # 0 = N/A
""", size=13)
text(s, 8.4, 1.75, 4.33, 4.9, [
    ("إزاي اكتشفت المشاكل؟", {"bold": True, "color": NAVY, "size": 20, "space_after": 10}),
    ("دالة quality_report بتعرض لكل عمود: النوع، القيم الناقصة، وعدد القيم المختلفة.", {"space_after": 8}),
    ("فحص التكرار: مفيش صفوف ولا id مكررة.", {"space_after": 8}),
    ("value_counts للأعمدة النصية كشف التسميات غير المتسقة.", {"space_after": 8}),
    ("describe وbox plots كشفوا القيم المتطرفة والتقييمات بصفر.", {}),
], size=15)
notes(s, "قبل أي تنظيف لازم أفهم البيانات. عملت دالة quality_report بتطلع ملخص لكل عمود. وبعدين فحصت التكرار، "
         "وعرضت القيم الموجودة في الأعمدة النصية، واستخدمت describe والـ box plots علشان أشوف القيم الغريبة.")

# ================================================================ 6. Problems
s = content_slide("المشاكل اللي لقيتها في البيانات", "الجزء الثالث — التقييم")
problems = [("83", "قيمة ناقصة في Arrival Delay ومتخزنة float"),
            ("4,202", "تقييم بصفر = «لا ينطبق» وليس أسوأ تقييم"),
            ("4", "أعمدة فيها تسميات غير متسقة مثل disloyal Customer و Eco"),
            ("1", "عمود Unnamed: 0 مالوش لازمة"),
            ("3,538", "قيمة متطرفة في التأخير — طلعت حقيقية"),
            ("14", "اسم عمود فيه مسافات وشرطات")]
for i, (num, desc) in enumerate(problems):
    col, row = i % 3, i // 3
    x, y = 8.85 - col * 4.12, 1.75 + row * 2.55
    card(s, x, y, 3.88, 2.25)
    text(s, x + 0.3, y + 0.25, 3.28, 0.75, num, size=36, color=ORANGE, bold=True, rtl=False, align=PP_ALIGN.RIGHT)
    text(s, x + 0.3, y + 1.1, 3.28, 1.0, desc, size=15, color=INK)
notes(s, "دي المشاكل اللي طلعت: 83 قيمة ناقصة في تأخير الوصول، وأهم مشكلة هي 4202 تقييم بصفر، والصفر هنا معناه "
         "إن السؤال ما ينطبقش على الراكب مش إنه أسوأ تقييم، ولو سبتها هتنزّل المتوسطات غلط. كمان في تسميات مش "
         "متسقة، وعمود index مالوش لازمة، وأسماء أعمدة صعبة في الكود، وقيم متطرفة في التأخير اتأكدت إنها حقيقية.")

# ================================================================ 7. Cleaning 1
s = content_slide("تنظيف الأعمدة والتسميات", "الجزء الرابع — التنظيف")
code_block(s, 0.6, 1.7, 7.4, 4.95, """
clean_df = raw_df.copy()
clean_df = clean_df.drop(columns=["Unnamed: 0"])

def to_snake_case(name):
    return (name.strip().lower()
            .replace("/", " ").replace("-", " ")
            .replace(" in minutes", "")
            .replace(" ", "_"))

def standardise_labels(df, mappings):
    for column, mapping in mappings.items():
        unmapped = set(df[column].unique()) - set(mapping)
        if unmapped:
            raise ValueError(unmapped)
        df[column] = df[column].map(mapping)
    return df
""", size=12)
text(s, 8.4, 1.75, 4.33, 4.9, [
    ("Define → Code → Test", {"bold": True, "color": NAVY, "size": 20, "space_after": 10, "rtl": False,
                              "align": PP_ALIGN.RIGHT}),
    ("بشتغل على نسخة copy علشان البيانات الأصلية تفضل زي ما هي.", {"space_after": 8}),
    ("to_snake_case: تحويل كل أسماء الأعمدة لشكل موحد سهل في الكود.", {"space_after": 8}),
    ("standardise_labels: توحيد التسميات، ولو في قيمة مش متغطية بيطلع Error — ده بيضمن إن مفيش حاجة اتنست.",
     {"space_after": 8}),
    ("بعد كل خطوة في Test بـ assert أو print.", {"color": MUTED}),
], size=15)
notes(s, "التنظيف ماشي بطريقة Define ثم Code ثم Test. بشتغل على نسخة من البيانات. حذفت العمود اللي مالوش لازمة، "
         "وعملت دالة تحوّل أسماء الأعمدة لـ snake case، ودالة توحد التسميات وبتطلع خطأ لو في قيمة نسيتها، "
         "وبعد كل خطوة بختبر النتيجة.")

# ================================================================ 8. Cleaning 2
s = content_slide("تنظيف القيم الناقصة والتقييمات بصفر", "الجزء الرابع — التنظيف")
code_block(s, 0.6, 1.7, 7.4, 4.95, """
# Missing arrival delay -> same passenger's departure delay
clean_df["arrival_delay"] = (clean_df["arrival_delay"]
    .fillna(clean_df["departure_delay"]).astype(int))

# Rating 0 = "Not Applicable" -> NaN
clean_df[service_columns] = (clean_df[service_columns]
    .replace(0, np.nan))

# New features
clean_df = add_features(clean_df)   # age / distance /
                                    # delay groups ...
clean_df.to_csv("Data_Cleaned.csv", index=False)
""", size=12)
text(s, 8.4, 1.75, 4.33, 4.9, [
    ("قرارات مهمة", {"bold": True, "color": NAVY, "size": 20, "space_after": 10}),
    ("ملء التأخير الناقص بتأخير الإقلاع لنفس الراكب، لأن الارتباط بينهم 0.96 — أدق من المتوسط أو الوسيط.",
     {"space_after": 8}),
    ("الصفر اتحول لـ NaN فبيتجاهل في المتوسطات.", {"space_after": 8}),
    ("القيم المتطرفة اتسابت لأنها حقيقية، واتقسمت لفئات.", {"space_after": 8}),
    ("أعمدة جديدة: age_group, distance_group, delay_group, average_service_rating, is_satisfied",
     {"color": MUTED, "size": 13}),
], size=15)
notes(s, "هنا أهم قرارات التنظيف. التأخير الناقص ملأته بتأخير الإقلاع لنفس الراكب لأن الاتنين مرتبطين جدًا. "
         "التقييمات بصفر حولتها لقيم فاضية علشان ما تأثرش على المتوسط. ما حذفتش القيم المتطرفة لأنها حقيقية. "
         "وضفت أعمدة جديدة زي الفئات العمرية وفئات المسافة والتأخير، وفي الآخر حفظت البيانات النظيفة في ملف.")

# ================================================================ 9. Univariate
s = content_slide("التحليل أحادي المتغير", "الجزء الخامس — التحليل الاستكشافي")
code_block(s, 0.6, 1.7, 6.6, 3.2, """
def plot_category_counts(df, column, ax):
    counts = df[column].value_counts()
    ax.bar(counts.index, counts.values)
    ...

for column, ax in zip(profile_columns, axes.flat):
    plot_category_counts(clean_df, column, ax)
""", size=12)
text(s, 0.6, 5.15, 6.6, 1.6, "دوال قابلة لإعادة الاستخدام بدل تكرار نفس الكود لكل عمود — وده من معايير التقييم.",
     size=15, color=MUTED)
univariate = [("الرضا", "Pie chart"), ("النوع، الولاء، نوع السفر، الدرجة", "Bar charts"),
              ("السن والمسافة", "Histogram + KDE"), ("التأخير", "Histogram (log scale)"),
              ("الخدمات الـ 14", "Bar + Heatmap")]
for i, (var, chart) in enumerate(univariate):
    y = 1.7 + i * 1.0
    card(s, 7.6, y, 5.13, 0.85)
    text(s, 9.9, y + 0.2, 2.6, 0.5, var, size=15, color=NAVY, bold=True)
    text(s, 7.8, y + 0.22, 2.1, 0.5, chart, size=13, color=SKY, bold=True, rtl=False)
notes(s, "في التحليل أحادي المتغير بدرس كل متغير لوحده. عملت دوال زي plot_category_counts وplot_numeric_distribution "
         "علشان ما أكررش الكود. استخدمت pie للرضا، وأعمدة للمتغيرات الفئوية، وهيستوجرام للسن والمسافة، "
         "ومقياس لوغاريتمي للتأخير لأن معظم القيم صفر.")

# ================================================================ 10. Bivariate & stats
s = content_slide("التحليل ثنائي ومتعدد المتغيرات + الإحصاء", "الجزء الخامس — التحليل الاستكشافي")
code_block(s, 0.6, 1.7, 7.4, 4.95, """
def chi_square_test(df, column):
    table = pd.crosstab(df[column], df["satisfaction"])
    chi2, p_value, _, _ = stats.chi2_contingency(table)
    cramers_v = np.sqrt(chi2 / (table.values.sum()
                                * (min(table.shape) - 1)))
    return p_value, cramers_v

def mann_whitney_test(df, column):
    satisfied = df.loc[df["satisfaction"] == "Satisfied",
                       column].dropna()
    dissatisfied = df.loc[df["satisfaction"] ==
                   "Neutral or Dissatisfied", column].dropna()
    return stats.mannwhitneyu(satisfied, dissatisfied)
""", size=12)
text(s, 8.4, 1.75, 4.33, 4.9, [
    ("كل نتيجة ليها اختبار", {"bold": True, "color": NAVY, "size": 20, "space_after": 10}),
    ([("Chi-square + Cramér's V", {"bold": True, "color": SKY}),
      (" للمتغيرات الفئوية مع الرضا", {})], {"space_after": 8}),
    ([("Mann-Whitney", {"bold": True, "color": SKY}), (" للمتغيرات الرقمية لأنها مش طبيعية التوزيع", {})],
     {"space_after": 8}),
    ([("Spearman", {"bold": True, "color": SKY}), (" لارتباط التقييمات بالرضا", {})], {"space_after": 12}),
    ("p-value أقل من 0.05 = فرق حقيقي مش صدفة.", {"color": MUTED}),
], size=15)
notes(s, "في التحليل الثنائي ومتعدد المتغيرات بربط كل متغير بالرضا. وأي استنتاج لازم يكون وراه اختبار إحصائي: "
         "Chi-square مع Cramér's V للمتغيرات الفئوية علشان أعرف قوة العلاقة، وMann-Whitney للمتغيرات الرقمية، "
         "وSpearman للتقييمات. لو p-value أقل من 0.05 يبقى الفرق حقيقي.")

# ================================================================ 11. Chart types
s = content_slide("أنواع الرسوم المستخدمة", "الجزء الخامس — التحليل الاستكشافي")
charts = ["Pie", "Bar / Count", "Stacked bar", "Histogram + KDE", "Box plot", "Violin", "Heatmap",
          "Scatter", "Line", "Dumbbell"]
for i, name in enumerate(charts):
    col, row = i % 5, i // 5
    x, y = 10.33 - col * 2.43, 1.8 + row * 1.5
    card(s, x, y, 2.25, 1.25)
    badge(s, x + 0.85, y + 0.12, str(i + 1), fill=SKY if i % 2 == 0 else NAVY, size=0.5, font_size=14)
    text(s, x + 0.1, y + 0.72, 2.05, 0.45, name, size=14, color=NAVY, bold=True, rtl=False, align=PP_ALIGN.CENTER)
text(s, 0.6, 5.1, 12.13, 1.6, [
    ([("10 أنواع رسوم", {"bold": True, "color": SKY}), (" (المطلوب 5 على الأقل)", {})], {"space_after": 8}),
    ([("أكثر من 10 متغيرات", {"bold": True, "color": SKY}),
      (" اتدرست أحادي وثنائي ومتعدد (المطلوب 6 على الأقل)", {})], {}),
], size=18)
notes(s, "استخدمت عشر أنواع رسوم مختلفة، والمطلوب خمسة على الأقل، ودرست أكتر من عشر متغيرات بالتحليل الأحادي "
         "والثنائي والمتعدد، والمطلوب ستة على الأقل.")

# ================================================================ 12. Results
s = content_slide("أهم النتائج", "الجزء السادس — النتائج")
results = [("43.9%", "بس من الركاب راضيين", SKY),
           ("70% و 19%", "رضا درجة البيزنس والسياحية", SKY),
           ("10%", "رضا المسافرين لأسباب شخصية — في كل الدرجات", ORANGE),
           ("+1.44", "فجوة الصعود الإلكتروني: أقوى خدمة مؤثرة", SKY),
           ("2.81 / 5", "الواي فاي أقل خدمة متقيّمة", ORANGE),
           ("من 48% لـ 35%", "الرضا من الوصول في الميعاد لتأخير أكتر من ساعة", ORANGE)]
for i, (num, lab, colr) in enumerate(results):
    col, row = i % 3, i // 3
    x, y = 8.85 - col * 4.12, 1.75 + row * 2.55
    card(s, x, y, 3.88, 2.25)
    text(s, x + 0.3, y + 0.3, 3.28, 0.8, num, size=32, color=colr, bold=True, align=PP_ALIGN.RIGHT)
    text(s, x + 0.3, y + 1.2, 3.28, 0.9, lab, size=15, color=INK)
notes(s, "أهم النتائج: 44 في المية بس راضيين. الدرجة ونوع السفر أقوى عاملين: البيزنس 70 في المية والسياحية 19. "
         "المسافرين لأسباب شخصية رضاهم حوالي 10 في المية في كل الدرجات. الصعود الإلكتروني أهم خدمة، والواي فاي "
         "أقل خدمة متقيمة رغم إنها مؤثرة جدًا. والتأخير بيقلل الرضا بس تأثيره ضعيف. والنوع ملوش أي تأثير.")

# ================================================================ 13. Conclusion
s = dark_slide()
text(s, 0.6, 0.7, 12.13, 0.5, "الخلاصة", size=18, color=LIGHT_MUTED, bold=True)
text(s, 0.6, 1.25, 12.13, 1.0, "الرضا بيتكسب في الكابينة وأونلاين — مش بالمسافة ولا بالنوع", size=34, color=WHITE,
     bold=True)
recs = [("1", "تحسين الرحلة الرقمية", "واي فاي، صعود وحجز إلكتروني"),
        ("2", "تطوير الدرجة السياحية", "المقاعد، الترفيه، المساحة"),
        ("3", "عروض للمسافرين لأسباب شخصية", "أقل شريحة رضا")]
for i, (n, head, sub) in enumerate(recs):
    x = 8.95 - i * 4.1
    shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, x, 3.0, 3.8, 2.3, NAVY_2)
    text(s, x + 0.3, 3.3, 3.2, 0.7, n, size=34, color=SKY, bold=True, rtl=False, align=PP_ALIGN.RIGHT)
    text(s, x + 0.3, 4.05, 3.2, 0.5, head, size=18, color=WHITE, bold=True)
    text(s, x + 0.3, 4.6, 3.2, 0.5, sub, size=14, color=LIGHT_MUTED)
text(s, 0.6, 5.9, 12.13, 0.9, "البيانات اتنظفت بالكامل، واتحللت من زوايا كتير، وكل نتيجة اتأكدت باختبار إحصائي. "
     "والبيانات النظيفة هي اللي بيستخدمها الداشبورد التفاعلي على Streamlit.", size=16, color=LIGHT_MUTED)
notes(s, "الخلاصة إن رضا الركاب بيعتمد على جودة الخدمة والتجربة الرقمية والدرجة، مش على المسافة ولا النوع. "
         "وأهم التوصيات: تحسين الواي فاي والصعود الإلكتروني، وتطوير الدرجة السياحية، وعروض للمسافرين لأسباب شخصية.")

# ================================================================ 14. Thank you
s = dark_slide()
shape(s, MSO_SHAPE.OVAL, -0.6, 1.2, 5.0, 5.0, NAVY_2)
shape(s, MSO_SHAPE.OVAL, 0.5, 4.4, 2.2, 2.2, SKY)
text(s, 4.7, 2.3, 8.0, 1.2, "شكرًا لحسن الاستماع", size=52, color=WHITE, bold=True)
text(s, 4.7, 3.7, 8.0, 0.6, "الكود والبيانات والداشبورد متاحين على GitHub", size=22, color=LIGHT_MUTED)
text(s, 4.7, 4.8, 8.0, 0.5, "github.com/hahmed1936/Mid-Project", size=18, color=WHITE, rtl=False,
     align=PP_ALIGN.RIGHT)
notes(s, "شكرًا لحسن الاستماع. كل الكود والبيانات والداشبورد موجودين على GitHub.")

prs.save(OUT)
print("saved", OUT, len(prs.slides), "slides")
